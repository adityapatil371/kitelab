"""Selling half early: does it matter WHERE you sell it?

    python -m scripts.scaleout_r_test

The scale-out rule banks half the position once the trade has made 1R -- one R
being the risk you took, entry minus stop. So on a Rs100 entry with a Rs90 stop,
R is Rs10 and half comes off at Rs110. This asks what happens if you move that
line: bank sooner at 0.5R, or hold out for 2R or 3R.

TRADE LEVEL ONLY, and that is not a stylistic choice. The one-account simulator
re-derives cash from shares x exit_price, and a scale-out has TWO exits at two
prices; the banked leg is folded into gross_profit and never reaches the trade
record. Verified 2026-08-31: of 1,523 banked trades in a 40-symbol sample, 1,522
disagree with the account engine's formula, one of them by enough to flip its
sign. So every number here is a sum over trades, each separately funded at a
fixed Rs1,00,000 book. Do not quote a CAGR off this sheet.
"""
from __future__ import annotations

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from kitelab import backtest, config, report, strategies

MULTIPLES = [0.5, 1.0, 1.5, 2.0, 3.0]
RUNS = [
    ("EMA M/W/D", "half", lambda s, r: backtest.simulate(s, scale_out="half", scale_r=r)),
    ("EMA M/W/D", "half_be",
     lambda s, r: backtest.simulate(s, scale_out="half_be", scale_r=r)),
    ("ATH Breakout", "half",
     lambda s, r: strategies.ath_breakout_trades(s, trailing_stops=True,
                                                 scale_out="half", scale_r=r)),
]
BASELINES = [
    ("EMA M/W/D", lambda s: backtest.simulate(s)),
    ("ATH Breakout", lambda s: strategies.ath_breakout_trades(s, trailing_stops=True)),
]


def collect(build, symbols) -> list[dict]:
    out = []
    for symbol in symbols:
        try:
            out.extend(build(symbol))
        except (SystemExit, FileNotFoundError):
            pass
    return out


def summarise(trades: list[dict]) -> dict:
    net = np.array([t["net_profit"] for t in trades], dtype=float)
    wins, losses = net[net > 0], net[net <= 0]
    banked = sum(1 for t in trades if "banked" in t["exit_reason"])
    gross_win, gross_loss = float(wins.sum()), float(-losses.sum())
    return {
        "trades": len(trades),
        "banked_pct": round(100 * banked / len(trades), 1) if trades else 0.0,
        "win_rate": round(100 * len(wins) / len(net), 1),
        "avg_win": round(float(wins.mean()) if len(wins) else 0.0),
        "avg_loss": round(float(losses.mean()) if len(losses) else 0.0),
        "expectancy": round(float(net.mean())),
        "net": round(float(net.sum())),
        "pf": round(gross_win / gross_loss, 2) if gross_loss else None,
    }


def main() -> None:
    symbols = config.load().all_symbols

    print("\n  baselines -- keep the whole position, never bank early:")
    base = {}
    for label, build in BASELINES:
        stats = summarise(collect(build, symbols))
        base[label] = stats
        print(f"    {label:<14} trades {stats['trades']:>5}  win {stats['win_rate']:>5}%"
              f"  expectancy {stats['expectancy']:>7,}  net {stats['net']:>12,}"
              f"  PF {stats['pf']}", flush=True)

    rows = []
    print(f"\n  {'strategy':<14}{'rule':<9}{'sell half at':>13}{'banked':>8}{'win%':>7}"
          f"{'expectancy':>12}{'net':>13}{'PF':>6}{'vs keeping all':>16}")
    print("  " + "-" * 100)
    for label, variant, build in RUNS:
        for multiple in MULTIPLES:
            stats = summarise(collect(lambda s, r=multiple: build(s, r), symbols))
            keep = base[label]["net"]
            stats.update(strategy=label, variant=variant, multiple=multiple,
                         vs_keep=round(100 * (stats["net"] - keep) / abs(keep), 1))
            rows.append(stats)
            print(f"  {label:<14}{variant:<9}{str(multiple) + 'R':>13}"
                  f"{stats['banked_pct']:>7}%{stats['win_rate']:>7}"
                  f"{stats['expectancy']:>12,}{stats['net']:>13,}{stats['pf']:>6}"
                  f"{stats['vs_keep']:>15}%", flush=True)

    book = Workbook()
    book.remove(book.active)
    sheet = book.create_sheet("Scale-out R")
    sheet["A1"] = (
        f"Banking half the position at a multiple of R, all "
        f"{len(config.load().all_symbols)} stocks, all history. "
        "R is the risk taken: entry minus the initial stop. TRADE LEVEL -- every "
        "signal separately funded at a fixed Rs1,00,000 book, summed. There is no "
        "CAGR here on purpose: the one-account simulator prices a trade as shares x "
        "one exit price, and a scale-out has two, so an account figure for these "
        "would be wrong. 'vs keeping all' compares to never banking early.")
    columns = [("Strategy", 15), ("Rule", 10), ("Sell half at", 13), ("Trades", 9),
               ("% that banked", 14), ("Win %", 9), ("Avg win", 11), ("Avg loss", 11),
               ("Expectancy", 12), ("Total net", 14), ("Profit factor", 13),
               ("vs keeping all %", 17)]
    for index, (name, width) in enumerate(columns, start=1):
        sheet.cell(3, index, name).font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    line = 4
    for label, _b in BASELINES:
        stats = base[label]
        values = [label, "keep all", "-", stats["trades"], 0.0, stats["win_rate"],
                  stats["avg_win"], stats["avg_loss"], stats["expectancy"],
                  stats["net"], stats["pf"], 0.0]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(line, index, value)
            cell.font = Font(bold=True)
            if index in (7, 8, 9, 10):
                cell.number_format = "#,##0"
        line += 1
    for stats in rows:
        values = [stats["strategy"], stats["variant"], f"{stats['multiple']:g}R",
                  stats["trades"], stats["banked_pct"], stats["win_rate"],
                  stats["avg_win"], stats["avg_loss"], stats["expectancy"],
                  stats["net"], stats["pf"], stats["vs_keep"]]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(line, index, value)
            if index in (7, 8, 9, 10):
                cell.number_format = "#,##0"
            if index in (5, 6, 12):
                cell.number_format = "0.0"
        line += 1
    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor="BDD7EE")

    print(f"\n  written: {report.save(book, 'Scale-Out R Sweep.xlsx')}")


if __name__ == "__main__":
    main()
