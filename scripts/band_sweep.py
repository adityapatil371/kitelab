"""Band sweep 0% -> 5% on all three EMA timeframe stacks.

    python -m scripts.band_sweep

The band is the dead zone around the EMAs: with band=B you must close B%
ABOVE all three EMAs to buy, and B% BELOW any one of them to be forced out.
Band 0% is the bare "close above/below the EMA" rule.

Writes output/EMA Band Sweep.xlsx in the same layout as the timeframe
comparison workbook: one sheet per stack, per-stock trade blocks with live
formulas (Shares / Cost / Profit / Cum Profit / Points), a total row under
each block (stock's own move, points captured), and the summary block on the
right (Win Trades / Total Win / Avg Win / Lose Trades / Total Loss / Avg Loss
/ Expectancy / Capture / Risk Reward) -- repeated for every band.

Five assigned stocks, class close-only convention, common window per stock
(hourly data starts 2015), fixed Rs1,00,000 capital and 1% risk per trade.
Trade sheets are GROSS; the Summary sheet also carries charges and net,
since avoiding charges is what a band is for.
"""
from __future__ import annotations

import math
import re
import shutil
import zipfile

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from kitelab import report
from scripts.tf_compare import ASSIGNED, simulate_variant, window_start

STACKS = [("QMW", "Q-M-W"), ("MWD", "M-W-D"), ("WDH", "W-D-H")]
BANDS = [0.0, 0.01, 0.02, 0.03, 0.04, 0.05]
TRADE_HEADERS = ["Stock", "Entry Date", "Exit Date", "Entry Price", "Stop Loss", "Exit",
                 "Exit Reason", "No. of Shares", "Cost of Entry", "Profit", "Cum Profit",
                 "Points"]
TRADE_WIDTHS = [11, 17, 17, 11, 11, 11, 16, 12, 13, 11, 12, 9]
SUM_HEADERS = ["Band", "Stock", "Win Trades", "Total Win", "Avg Win", "Lose Trades",
               "Total Loss", "Avg Loss", "Expectancy", "Capture", "Risk Reward"]
BAND_FILL = PatternFill("solid", fgColor="BDD7EE")


def collect(variant: str, band: float) -> dict[str, list[dict]]:
    out = {}
    for symbol in ASSIGNED:
        start = window_start(symbol)
        trades = [t for t in simulate_variant(symbol, variant, band=band)
                  if pd.Timestamp(t["entry_ts"]).normalize() >= start]
        out[symbol] = sorted(trades, key=lambda t: t["entry_ts"], reverse=True)
    return out


def main() -> None:
    data = {(v, b): collect(v, b) for v, _ in STACKS for b in BANDS}
    for (v, b), per_stock in data.items():
        flat = [t for lst in per_stock.values() for t in lst]
        gross = sum(t["gross_profit"] for t in flat)
        charges = sum(t["charges"] for t in flat)
        print(f"  {v:<4} band {b:.0%}  {len(flat):>5} trades  gross {gross:>10,.0f}  "
              f"charges {charges:>9,.0f}  net {gross-charges:>10,.0f}", flush=True)

    book = Workbook()
    book.remove(book.active)
    cached: dict[tuple[str, str], float] = {}

    # ---------- Summary sheet ----------
    s_sheet = book.create_sheet("Summary")
    s_sheet["A1"] = ("Band sweep 0%-5%: the dead zone around the EMAs, on all three stacks. "
                     "Five assigned stocks, common window each, class close-only convention. "
                     "GROSS is before charges, NET after Zerodha delivery charges -- a band "
                     "earns its keep by cutting trade count, so judge it on NET.")
    s_cols = ["Stack", "Band", "Trades", "Win %", "Avg Win", "Avg Loss", "Expectancy",
              "Profit Factor", "Gross", "Charges", "Net", "Charges % of Gross",
              "Median Hold (days)", "Median Stop %"]
    s_fmts = ["@", "0%", "0", "0.0%", "#,##0", "#,##0", "#,##0", "0.00", "#,##0", "#,##0",
              "#,##0", "0.0", "0", "0.00"]
    for c, (name, width) in enumerate(zip(s_cols, [9, 7, 8, 8, 10, 10, 11, 8, 12, 10, 12, 11, 12, 11]), start=1):
        s_sheet.cell(3, c, name).font = Font(bold=True)
        s_sheet.column_dimensions[get_column_letter(c)].width = width
    srow = 4
    for variant, label in STACKS:
        for band in BANDS:
            flat = [t for lst in data[(variant, band)].values() for t in lst]
            g = [t["gross_profit"] for t in flat]
            wins = [x for x in g if x > 0]
            losses = [x for x in g if x <= 0]
            charges = sum(t["charges"] for t in flat)
            vals = [label, band, len(flat), len(wins) / len(g) if g else 0,
                    round(np.mean(wins)) if wins else 0, round(np.mean(losses)) if losses else 0,
                    round(sum(g) / len(g)) if g else 0,
                    round(sum(wins) / -sum(losses), 2) if losses and sum(losses) < 0 else None,
                    round(sum(g)), round(charges), round(sum(g) - charges),
                    round(100 * charges / sum(g), 1) if sum(g) else None,
                    round(float(np.median([t["days_held"] for t in flat]))) if flat else 0,
                    round(float(np.median([t["stop_pct"] for t in flat])), 2) if flat else 0]
            for c, (val, fmt) in enumerate(zip(vals, s_fmts), start=1):
                cell = s_sheet.cell(srow, c, val)
                cell.number_format = fmt
            srow += 1
        srow += 1

    # ---------- one sheet per stack ----------
    for variant, label in STACKS:
        ws = book.create_sheet(label)
        ws["A1"] = "Capital"; ws["B1"] = 100000
        ws["C1"] = "Risk %";  ws["D1"] = 0.01
        for a in ("A1", "C1"):
            ws[a].font = Font(bold=True)
        for c, (name, width) in enumerate(zip(TRADE_HEADERS, TRADE_WIDTHS), start=1):
            ws.cell(3, c, name).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(c)].width = width
        for c, name in enumerate(SUM_HEADERS, start=14):          # N onwards
            ws.cell(3, c, name).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(c)].width = 12
        date_fmt = "yyyy-mm-dd hh:mm" if variant == "WDH" else "yyyy-mm-dd"

        row = 4
        sum_row = 4
        for band in BANDS:
            marker = ws.cell(row, 1, f"BAND {band:.0%}")
            marker.font = Font(bold=True)
            marker.fill = BAND_FILL
            row += 1
            for symbol in ASSIGNED:
                trades = data[(variant, band)][symbol]
                if not trades:
                    continue
                lo = row
                cum = 0.0
                for t in trades:
                    d = round(t["entry_price"], 2)
                    e = round(t["stop"], 2)
                    f = round(t["exit_price"], 2)
                    ws.cell(row, 1, symbol)
                    ws.cell(row, 2, pd.Timestamp(t["entry_ts"]).to_pydatetime()).number_format = date_fmt
                    ws.cell(row, 3, pd.Timestamp(t["exit_ts"]).to_pydatetime()).number_format = date_fmt
                    for col, v in ((4, d), (5, e), (6, f)):
                        ws.cell(row, col, v).number_format = "0.00"
                    ws.cell(row, 7, t["exit_reason"])
                    ws.cell(row, 8, f"=MIN(ROUNDDOWN($B$1*$D$1/(D{row}-E{row}),0),ROUNDDOWN($B$1/D{row},0))").number_format = "0"
                    ws.cell(row, 9, f"=D{row}*H{row}").number_format = "#,##0.00"
                    ws.cell(row, 10, f"=(F{row}-D{row})*H{row}").number_format = "#,##0.00"
                    ws.cell(row, 11, f"=J{row}" if row == lo else f"=K{row-1}+J{row}").number_format = "#,##0.00"
                    ws.cell(row, 12, f"=F{row}-D{row}").number_format = "0.00"
                    shares = min(math.floor(100000 * 0.01 / (d - e)), math.floor(100000 / d))
                    profit = (f - d) * shares
                    cum += profit
                    cached[(label, f"H{row}")] = shares
                    cached[(label, f"I{row}")] = round(d * shares, 2)
                    cached[(label, f"J{row}")] = round(profit, 2)
                    cached[(label, f"K{row}")] = round(cum, 2)
                    cached[(label, f"L{row}")] = round(f - d, 2)
                    row += 1
                hi = row - 1
                # total row: the stock's own move, and the points the strategy captured
                ws.cell(row, 6, f"=F{lo}-D{hi}").number_format = "0.00"
                ws.cell(row, 12, f"=SUM(L{lo}:L{hi})").number_format = "0.00"
                move = round(trades[0]["exit_price"] - trades[-1]["entry_price"], 2)
                points = round(sum(round(t["exit_price"], 2) - round(t["entry_price"], 2)
                                   for t in trades), 2)
                cached[(label, f"F{row}")] = move
                cached[(label, f"L{row}")] = points
                total_row = row
                row += 2

                # summary block row for this (band, stock)
                profits = [cached[(label, f"J{r}")] for r in range(lo, hi + 1)]
                wins = [x for x in profits if x > 0]
                losses = [x for x in profits if x < 0]
                rng = f"J{lo}:J{hi}"
                ws.cell(sum_row, 14, f"{band:.0%}")
                ws.cell(sum_row, 15, symbol)
                ws.cell(sum_row, 16, f'=COUNTIF({rng},">0")')
                ws.cell(sum_row, 17, f'=SUMIF({rng},">0")').number_format = "#,##0.00"
                ws.cell(sum_row, 18, f'=IFERROR(SUMIF({rng},">0")/COUNTIF({rng},">0"),"-")').number_format = "#,##0.00"
                ws.cell(sum_row, 19, f'=COUNTIF({rng},"<0")')
                ws.cell(sum_row, 20, f'=SUMIF({rng},"<0")').number_format = "#,##0.00"
                ws.cell(sum_row, 21, f'=IFERROR(SUMIF({rng},"<0")/COUNTIF({rng},"<0"),"-")').number_format = "#,##0.00"
                ws.cell(sum_row, 22, f'=IFERROR((P{sum_row}/(P{sum_row}+S{sum_row}))*R{sum_row}'
                                     f'+(S{sum_row}/(P{sum_row}+S{sum_row}))*U{sum_row},"-")').number_format = "#,##0.00"
                ws.cell(sum_row, 23, f'=IFERROR(L{total_row}/F{total_row},"-")').number_format = "0.00"
                ws.cell(sum_row, 24, f'=IFERROR(R{sum_row}/(U{sum_row}*-1),"-")').number_format = "0.00"
                cached[(label, f"P{sum_row}")] = len(wins)
                cached[(label, f"Q{sum_row}")] = round(sum(wins), 2)
                cached[(label, f"S{sum_row}")] = len(losses)
                cached[(label, f"T{sum_row}")] = round(sum(losses), 2)
                if wins:
                    cached[(label, f"R{sum_row}")] = round(sum(wins) / len(wins), 2)
                if losses:
                    cached[(label, f"U{sum_row}")] = round(sum(losses) / len(losses), 2)
                if wins and losses:
                    n = len(wins) + len(losses)
                    exp = (len(wins) / n) * (sum(wins) / len(wins)) + (len(losses) / n) * (sum(losses) / len(losses))
                    cached[(label, f"V{sum_row}")] = round(exp, 2)
                    cached[(label, f"X{sum_row}")] = round((sum(wins) / len(wins)) / abs(sum(losses) / len(losses)), 2)
                if move:
                    cached[(label, f"W{sum_row}")] = round(points / move, 2)
                sum_row += 1
            sum_row += 1

    target = report.save(book, "EMA Band Sweep.xlsx")

    names = {f"xl/worksheets/sheet{i}.xml": n for i, n in enumerate(book.sheetnames, start=1)}
    tmp = str(target) + ".tmp"
    with zipfile.ZipFile(target) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            blob = zin.read(item)
            if item in names:
                name = names[item]
                xml = blob.decode()

                def fill(m, name=name):
                    v = cached.get((name, m.group(1)))
                    return m.group(0) if v is None else m.group(0).replace("<v />", f"<v>{v}</v>")

                xml, n = re.subn(r'<c r="([A-Z]+\d+)"[^>]*><f>[^<]*</f><v /></c>', fill, xml)
                if n:
                    print(f"    {name}: {n} formula cells given cached values")
                blob = xml.encode()
            zout.writestr(item, blob)
    shutil.move(tmp, target)
    print(f"\n  written: {target}")


if __name__ == "__main__":
    main()
