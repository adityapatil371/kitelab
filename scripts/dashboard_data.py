"""Precompute everything the local dashboard can show.

    python -m scripts.dashboard_data

Writes data/dashboard.json for http://localhost:8765/dashboard (served by
python -m scripts.chart). The page is a pure viewer -- every control selects
among these precomputed results, nothing is simulated in the browser.

Sections:
    grid        portfolio simulations: strategy (EMA class rule / Breakout)
                x universe (all/in/holdout) x risk (0.25-2%) x capital (50k-5L)
    scaleout    the teacher's "sell half at +1R" idea (and the breakeven
                variant), measured the way scaleout_test did: fixed capital,
                sum of trades, net of charges -- portfolio.run cannot price a
                two-part exit, so this deliberately stays trade-level
    stocks      every one of the 199 stocks: weekly closes + every EMA and
                Breakout trade + per-stock stats
    assets      the six class instruments: closes, trades, account curve,
                buy-and-hold comparison
    timeframes  Q/M/W vs M/W/D vs W/D/H on the five assigned stocks
    btc_validation   the manual Bitcoin backtest validation rows
    nifty       NIFTY 50 overlay series
"""
from __future__ import annotations

import json
import pickle
from pathlib import Path

import pandas as pd

from kitelab import backtest, config, frames, portfolio, sizing, strategies
from scripts.drawdown_report import bh_stats, episodes, underwater_stats
from scripts.tf_compare import (ASSIGNED, VARIANTS as TF_VARIANTS, simulate_variant,
                                summarise as tf_summarise, window_start)

CACHE = Path(__file__).resolve().parent.parent / "data" / "signal_cache"
OUT = Path(__file__).resolve().parent.parent / "data" / "dashboard.json"
BTC_VALIDATED = Path(__file__).resolve().parent.parent / "output" / "Aditya P bitcoin - validated.xlsx"
RISKS = [0.25, 0.5, 1.0, 2.0]
CAPITALS = [50_000, 100_000, 250_000, 500_000]
ASSETS = [("BITCOIN", "30m", 0.0010), ("NIFTY 50", "30m", 0.0005),
          ("NIFTY BANK", "30m", 0.0005), ("GOLD", "1d", 0.0005),
          ("SILVER", "1d", 0.0005), ("CRUDEOIL", "1d", 0.0005)]
STEP = 5   # keep every 5th trading day of curves (~weekly)


# ------------------------------------------------------------ helpers ----

def curve_payload(curve):
    days, eq, dd = [], [], []
    peak = float("-inf")
    for i, (day, equity) in enumerate(curve):
        peak = max(peak, equity)
        if i % STEP and i != len(curve) - 1:
            continue
        days.append(day.strftime("%Y-%m-%d"))
        eq.append(round(equity))
        dd.append(round(100 * (equity - peak) / peak, 1))
    return {"d": days, "eq": eq, "dd": dd}


def run_payload(r):
    longest, current = underwater_stats(r["curve"])
    eps = [{"peak": str(e["peak_day"].date()), "trough": str(e["trough_day"].date()),
            "depth": round(e["depth_pct"], 1),
            "recovered": str(e["recovered"].date()) if e["recovered"] is not None else None,
            "years": round(e["days"] / 365.25, 1)} for e in episodes(r["curve"])]
    return {"cagr": round(r["cagr_pct"], 1), "final": round(r["final"]),
            "ret": round(r["return_pct"]), "maxdd": round(r["max_drawdown_pct"], 1),
            "uw_long": round(longest / 365.25, 1), "uw_now": round(current / 365.25, 1),
            "taken": len(r["taken"]), "signals": r["signals"],
            "episodes": eps, "curve": curve_payload(r["curve"])}


def cached_signals(name: str, build) -> list[dict]:
    path = CACHE / f"{name}.pkl"
    if path.exists():
        return pickle.loads(path.read_bytes())
    cfg = config.load()
    out = []
    for index, symbol in enumerate(cfg.all_symbols, 1):
        try:
            out.extend(build(symbol))
        except SystemExit:
            pass
        if index % 25 == 0:
            print(f"    {name}: {index}/{len(cfg.all_symbols)}", flush=True)
    path.write_bytes(pickle.dumps(out))
    return out


def trade_stats(trades: list[dict]) -> dict:
    nets = [t["net_profit"] for t in trades]
    wins = [n for n in nets if n > 0]
    losses = [n for n in nets if n <= 0]
    running = peak = 0.0
    worst_run = 0.0
    for t in sorted(trades, key=lambda t: t["exit_ts"]):
        running += t["net_profit"]
        peak = max(peak, running)
        worst_run = min(worst_run, running - peak)
    return {"trades": len(trades), "wins": len(wins),
            "win_rate": round(len(wins) / len(nets), 3) if nets else 0,
            "net": round(sum(nets)),
            "pf": round(sum(wins) / -sum(losses), 2) if losses and sum(losses) < 0 else None,
            "best": round(max(nets)) if nets else 0,
            "worst": round(min(nets)) if nets else 0,
            "worst_run": round(worst_run)}


def slim_trades(trades: list[dict]) -> list[dict]:
    return [{"e": pd.Timestamp(t["entry_ts"]).strftime("%Y-%m-%d"),
             "x": pd.Timestamp(t["exit_ts"]).strftime("%Y-%m-%d"),
             "ep": round(t["entry_price"], 2), "xp": round(t["exit_price"], 2),
             "st": round(t["stop"], 2), "sh": round(t["shares"], 4),
             "net": round(t["net_profit"]), "why": t["exit_reason"]}
            for t in sorted(trades, key=lambda t: t["entry_ts"], reverse=True)]


def close_series(daily: pd.DataFrame) -> dict:
    d, c = [], []
    for i, rec in enumerate(daily.itertuples(index=False)):
        if i % STEP and i != len(daily) - 1:
            continue
        d.append(rec.ts.strftime("%Y-%m-%d"))
        c.append(round(float(rec.close), 2))
    return {"d": d, "c": c}


# --------------------------------------------------------------- main ----

def main() -> None:
    cfg = config.load()

    print("  signal lists (cached where possible):", flush=True)
    base = {
        "ema": cached_signals("EMA_199", backtest.simulate),
        "brk": cached_signals("Breakout_199",
                              lambda s: strategies.ath_breakout_trades(s, trailing_stops=True)),
    }
    scale_lists = {
        ("ema", "half"): cached_signals("EMA_half_199",
                                        lambda s: backtest.simulate(s, scale_out="half")),
        ("ema", "half_be"): cached_signals("EMA_halfbe_199",
                                           lambda s: backtest.simulate(s, scale_out="half_be")),
        ("brk", "half"): cached_signals("Breakout_half_199",
                                        lambda s: strategies.ath_breakout_trades(
                                            s, trailing_stops=True, scale_out="half")),
        ("brk", "half_be"): cached_signals("Breakout_halfbe_199",
                                           lambda s: strategies.ath_breakout_trades(
                                               s, trailing_stops=True, scale_out="half_be")),
    }

    universes = {"all": ("All 199 stocks", None),
                 "in": ("49 in-sample", set(cfg.in_sample)),
                 "out": ("150 holdout", set(cfg.out_of_sample))}

    grid = {}
    for skey, signals in base.items():
        for ukey, (ulabel, members) in universes.items():
            subset = (signals if members is None
                      else [t for t in signals if t["symbol"] in members])
            for risk in RISKS:
                for capital in CAPITALS:
                    r = portfolio.run(subset, capital, risk / 100)
                    grid[f"{skey}|{ukey}|{risk:g}|{capital}"] = run_payload(r)
            print(f"  grid: {skey} / {ulabel} done", flush=True)

    scaleout = {}
    for skey in ("ema", "brk"):
        rows = [{"variant": "Keep full position (baseline)", **trade_stats(base[skey])},
                {"variant": "Sell half at +1R", **trade_stats(scale_lists[(skey, "half")])},
                {"variant": "Sell half at +1R, stop to breakeven",
                 **trade_stats(scale_lists[(skey, "half_be")])}]
        scaleout[skey] = rows
        print(f"  scale-out: {skey} done", flush=True)

    print("  per-stock detail:", flush=True)
    by_symbol = {"ema": {}, "brk": {}}
    for skey, signals in base.items():
        for t in signals:
            by_symbol[skey].setdefault(t["symbol"], []).append(t)
    stocks = {}
    for index, symbol in enumerate(sorted(cfg.all_symbols), 1):
        try:
            daily = frames.daily(symbol)
        except SystemExit:
            continue
        ema_tr = by_symbol["ema"].get(symbol, [])
        brk_tr = by_symbol["brk"].get(symbol, [])
        stocks[symbol] = {
            "closes": close_series(daily),
            "assigned": symbol in ASSIGNED_SET,
            "ema": {"stats": trade_stats(ema_tr), "trades": slim_trades(ema_tr)},
            "brk": {"stats": trade_stats(brk_tr), "trades": slim_trades(brk_tr)},
        }
        if index % 50 == 0:
            print(f"    stocks: {index}/{len(cfg.all_symbols)}", flush=True)

    print("  assets detail:", flush=True)
    assets = {}
    for symbol, brk_tf, fee in ASSETS:
        backtest.FLAT_FEE_RATE = fee
        sizing.FRACTIONAL = True
        try:
            ema_tr = backtest.simulate(symbol)
            brk_tr = strategies.ath_breakout_trades(symbol, True, timeframe=brk_tf)
            daily = frames.daily(symbol)
            bh = bh_stats(daily)
            entry = {"closes": close_series(daily),
                     "bh": {"cagr": round(bh["cagr"], 1), "maxdd": round(bh["maxdd"], 1),
                            "uw": round(bh["longest_uw"] / 365.25, 1),
                            "years": round(bh["years"], 1)}}
            for strat, trades in (("ema", ema_tr), ("brk", brk_tr)):
                r = portfolio.run(trades, 100_000, 0.01) if trades else None
                entry[strat] = {"stats": trade_stats(trades), "trades": slim_trades(trades),
                                "account": run_payload(r) if r else None}
            assets[symbol] = entry
        finally:
            backtest.FLAT_FEE_RATE = None
            sizing.FRACTIONAL = False
        print(f"    {symbol} done", flush=True)

    tf = {}
    for key, label, desc in TF_VARIANTS:
        rows, allt = [], []
        for symbol in ASSIGNED:
            w = window_start(symbol)
            kept = [t for t in simulate_variant(symbol, key)
                    if pd.Timestamp(t["entry_ts"]).normalize() >= w]
            allt += kept
            s = tf_summarise(kept)
            rows.append({"sym": symbol, **{k: round(v, 2) if isinstance(v, float) else v
                        for k, v in s.items() if k != "charges"}})
        s = tf_summarise(allt)
        rows.append({"sym": "ALL 5", **{k: round(v, 2) if isinstance(v, float) else v
                     for k, v in s.items() if k != "charges"}})
        tf[key] = {"label": label, "desc": desc, "rows": rows}
        print(f"  timeframes: {label} done", flush=True)

    btc_rows = []
    if BTC_VALIDATED.exists():
        from openpyxl import load_workbook
        ws = load_workbook(BTC_VALIDATED, data_only=True).worksheets[0]
        for r in range(4, 40):
            if ws.cell(r, 1).value is None:
                continue
            btc_rows.append({
                "date": str(ws.cell(r, 1).value)[:10], "entry": ws.cell(r, 2).value,
                "stop": ws.cell(r, 3).value, "shares": ws.cell(r, 4).value,
                "claimed": ws.cell(r, 5).value, "reason": ws.cell(r, 6).value,
                "verdict": ws.cell(r, 8).value, "corrected": ws.cell(r, 9).value,
                "profit": ws.cell(r, 11).value})

    nifty = frames.daily("NIFTY 50")
    payload = {
        "built": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M"),
        "strategies": {"ema": "EMA 20 stack (class rule)", "brk": "ATH Breakout"},
        "universes": {k: v[0] for k, v in universes.items()},
        "risks": RISKS, "capitals": CAPITALS,
        "assigned": list(ASSIGNED),
        "grid": grid, "scaleout": scaleout, "stocks": stocks, "assets": assets,
        "timeframes": tf, "btc_validation": btc_rows,
        "nifty": close_series(nifty),
    }
    OUT.write_text(json.dumps(payload))
    print(f"\n  written: {OUT} ({OUT.stat().st_size/1e6:.1f} MB)")


ASSIGNED_SET = set(ASSIGNED)

if __name__ == "__main__":
    main()
