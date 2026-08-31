"""What execution really costs, measured instead of guessed.

    python -m scripts.slippage_report            # the scenarios and the workbook
    python -m scripts.slippage_report --audit    # only the Corwin-Schultz rejection

Until now every result in the lab assumed a perfect fill. The tornado study priced
that assumption at 4.8 CAGR points using a flat percentage haircut applied after the
fact. This replaces the flat guess with a per-stock, per-day model (kitelab.slippage)
and, separately, with the honest fill convention -- an end-of-day trader reads the
close after the bell and deals at the next open, not at the print they just watched.

Every scenario is the same reference account: EMA M/W/D, 2% band, the whole universe,
1% risk, Rs2,50,000.
"""
from __future__ import annotations

import argparse
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from kitelab import backtest, config, frames, portfolio, report, slippage

CAPITAL = 250_000.0
RISK = 0.01
K_CS = 3 - 2 * np.sqrt(2)


def corwin_schultz(day: pd.DataFrame) -> np.ndarray:
    """The textbook high-low spread estimator, kept so its rejection is reproducible."""
    high = day["high"].to_numpy(float)
    low = day["low"].to_numpy(float)
    high = np.where(high > 0, high, np.nan)
    low = np.where(low > 0, low, np.nan)
    two_high = np.maximum(high[1:], high[:-1])
    two_low = np.minimum(low[1:], low[:-1])
    beta = np.log(high[:-1] / low[:-1]) ** 2 + np.log(high[1:] / low[1:]) ** 2
    gamma = np.log(two_high / two_low) ** 2
    alpha = (np.sqrt(2 * beta) - np.sqrt(beta)) / K_CS - np.sqrt(gamma / K_CS)
    return 2 * (np.exp(alpha) - 1) / (1 + np.exp(alpha))


def audit(symbols) -> pd.DataFrame:
    """Why we do not use Corwin-Schultz: it measures volatility, not spread."""
    rows = []
    for symbol in symbols:
        try:
            day = frames.daily(symbol)
        except SystemExit:
            continue
        if len(day) < 300:
            continue
        spread = corwin_schultz(day)
        usable = np.isfinite(spread)
        clipped = np.where(usable & (spread > 0), spread, 0.0)
        turnover = (day["close"] * day["volume"]).to_numpy(float)
        rows.append({
            "symbol": symbol,
            "bars": len(day),
            "adv_cr": float(np.nanmedian(turnover)) / 1e7,
            "cs_half_bps": float(np.nanmean(clipped)) / 2 * 10_000,
            "cs_negative_frac": float(np.mean(usable & (spread <= 0))),
        })
    frame = pd.DataFrame(rows).sort_values("adv_cr").reset_index(drop=True)
    frame["ladder_half_bps"] = [
        slippage.half_spread(r.symbol, frames.daily(r.symbol)["ts"].iloc[-1],
                             float(frames.daily(r.symbol)["close"].iloc[-1])) * 10_000
        for r in frame.itertuples()]
    return frame


def run_scenario(symbols, *, enabled: bool, next_open: bool,
                 spread_k: float | None = None, impact_c: float | None = None,
                 participation: float | None = None) -> dict:
    """One full pass over the universe under one execution model."""
    slippage.ENABLED = enabled
    backtest.NEXT_OPEN_FILLS = next_open
    slippage.MAX_PARTICIPATION = participation
    slippage.SPREAD_K = 30.0 if spread_k is None else spread_k
    slippage.IMPACT_C = 0.5 if impact_c is None else impact_c
    slippage.reset()

    trades = []
    for symbol in symbols:
        try:
            trades.extend(backtest.simulate(symbol))
        except SystemExit:
            pass
    result = portfolio.run(trades, CAPITAL, RISK)
    taken = result["taken"]
    slipped = sum(abs(t["entry_price"] - t.get("quoted_entry", t["entry_price"])) * t["shares"]
                  + abs(t.get("quoted_exit", t["exit_price"]) - t["exit_price"]) * t["shares"]
                  for t in taken)
    return {
        "signals": len(trades),
        "taken": len(taken),
        "cagr": report.cagr_cell(result),
        "maxdd": result["max_drawdown_pct"],
        "final": result["final"],
        "slippage_rs": slipped,
        "skipped_liquidity": result["skipped_liquidity"],
        "charges_rs": sum(backtest.charges(t["shares"] * t["entry_price"],
                                           t["shares"] * t["exit_price"],
                                           t["same_session"]) for t in taken),
    }


def restore() -> None:
    slippage.ENABLED = False
    backtest.NEXT_OPEN_FILLS = False
    slippage.SPREAD_K = 30.0
    slippage.IMPACT_C = 0.5
    slippage.MAX_PARTICIPATION = None
    slippage.reset()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit", action="store_true",
                        help="only reproduce the Corwin-Schultz rejection")
    args = parser.parse_args()
    symbols = config.load().all_symbols

    table = audit(symbols)
    print(f"\n  Corwin-Schultz audit over {len(table)} stocks with usable history")
    print(f"    CS half-spread   median {table.cs_half_bps.median():.1f} bps"
          f"   min {table.cs_half_bps.min():.1f}   max {table.cs_half_bps.max():.1f}")
    print(f"    most liquid name {table.iloc[-1].symbol} "
          f"(Rs{table.iloc[-1].adv_cr:.0f} cr/day) -> CS says "
          f"{table.iloc[-1].cs_half_bps:.1f} bps, ladder says "
          f"{table.iloc[-1].ladder_half_bps:.1f} bps")
    print(f"    thinnest stock's spread as a multiple of the most liquid one's: "
          f"CS says {table.cs_half_bps.iloc[0] / table.cs_half_bps.iloc[-1]:.1f}x, "
          f"ladder says {table.ladder_half_bps.iloc[0] / table.ladder_half_bps.iloc[-1]:.0f}x")
    print(f"    day-pairs returning a NEGATIVE spread: "
          f"{table.cs_negative_frac.mean():.0%}")
    if args.audit:
        return

    # Two different things get mixed up if you are not careful: what execution
    # COSTS, and what a position-size limit DOES to the strategy. The limit is not a
    # cost -- it is a sizing rule, and on this history it helps. So the cost of
    # execution is read DOWN the block that shares a sizing rule, not against the
    # old unfillable baseline.
    scenarios = [
        ("1  Old baseline: perfect fills, no size limit",
         dict(enabled=False, next_open=False)),
        ("2  Spread only, no size limit",
         dict(enabled=True, next_open=False, impact_c=0.0)),
        ("3  Impact only, no size limit",
         dict(enabled=True, next_open=False, spread_k=0.0)),
        ("4  Spread + impact, no size limit (unfillable)",
         dict(enabled=True, next_open=False)),
        ("5  Size limit only: order <= 1% of a day's turnover",
         dict(enabled=True, next_open=False, spread_k=0.0, impact_c=0.0,
              participation=0.01)),
        ("6  Size limit + spread",
         dict(enabled=True, next_open=False, impact_c=0.0, participation=0.01)),
        ("7  Size limit + spread + impact  <- the honest number",
         dict(enabled=True, next_open=False, participation=0.01)),
        ("8  ...also dealing at the next open, not the close",
         dict(enabled=True, next_open=True, participation=0.01)),
        ("9  As 7, spreads HALVED (K=15)",
         dict(enabled=True, next_open=False, participation=0.01, spread_k=15.0)),
        ("10 As 7, spreads DOUBLED (K=60)",
         dict(enabled=True, next_open=False, participation=0.01, spread_k=60.0)),
    ]

    print(f"\n  {'scenario':<50}{'CAGR':>8}{'maxDD':>9}{'taken':>8}"
          f"{'no liq':>8}{'slippage Rs':>14}")
    print("  " + "-" * 97)
    results = []
    for label, kwargs in scenarios:
        slippage.SPREAD_K, slippage.IMPACT_C = 30.0, 0.5
        out = run_scenario(symbols, **kwargs)
        results.append((label, out))
        print(f"  {label:<50}{out['cagr']:>7.1f}%{out['maxdd']:>8.1f}%"
              f"{out['taken']:>8}{out['skipped_liquidity']:>8}"
              f"{out['slippage_rs']:>14,.0f}", flush=True)
    restore()

    base = results[0][1]["cagr"]
    book = Workbook()
    book.remove(book.active)

    sheet = book.create_sheet("Scenarios")
    sheet["A1"] = (
        f"Execution cost on the reference account: EMA M/W/D, 2% band, "
        f"all {len(config.load().all_symbols)} stocks, "
        f"1% risk, Rs2,50,000. Row 1 is the old baseline, {base:.1f}% CAGR, which assumed "
        "a perfect fill at any size. Rows 2-4 add cost but keep that impossible sizing, "
        "so they overstate the damage: without a size limit this account puts up to 37x "
        "a stock's ENTIRE daily turnover into one position, and the impact model duly "
        "charges it a fortune for an order it could never place. Rows 5-7 are the ones "
        "to read: they share one realistic sizing rule, so the drop from row 5 to row 7 "
        "is what execution actually costs. Note the spread itself is ASSUMED -- Kite "
        "serves candles, not quotes, so a bid-ask cannot be measured from our data; "
        "rows 9 and 10 halve and double it. Impact, the size limit and the next-open "
        "fills are all driven by measured data.")
    columns = [("Scenario", 50), ("CAGR %", 10), ("vs old baseline (pts)", 21),
               ("COST: vs its own no-cost row (pts)", 33),
               ("Max drawdown %", 15), ("Trades taken", 13),
               ("Skipped: too illiquid", 20),
               ("Paid in slippage (Rs)", 20), ("Paid in charges (Rs)", 19),
               ("Final equity (Rs)", 18)]
    for index, (name, width) in enumerate(columns, start=1):
        sheet.cell(3, index, name).font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    for offset, (label, out) in enumerate(results):
        row = 4 + offset
        # Rows 1-4 share the old (unfillable) sizing; rows 5-10 share the 1% limit.
        # Comparing across those two blocks measures the sizing change, not the cost.
        reference = base if offset < 4 else results[4][1]["cagr"]
        values = [label, round(out["cagr"], 2), round(out["cagr"] - base, 2),
                  round(out["cagr"] - reference, 2),
                  round(out["maxdd"], 1), out["taken"], out["skipped_liquidity"],
                  round(out["slippage_rs"]), round(out["charges_rs"]),
                  round(out["final"])]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row, index, value)
            if index in (2, 3, 4, 5):
                cell.number_format = "0.0"
            if index in (8, 9, 10):
                cell.number_format = "#,##0"
        sheet.cell(row, 4).font = Font(bold=True)

    ladder = book.create_sheet("Spread ladder")
    ladder["A1"] = ("The assumed half-spread per stock, thinnest first. ADV is the median "
                    "daily traded value over the whole history. The ladder is "
                    "SPREAD_K/sqrt(ADV in crore) basis points, floored at half a tick.")
    head = [("Stock", 14), ("Median turnover (Rs cr/day)", 26), ("Ladder half-spread (bps)", 24),
            ("Corwin-Schultz says (bps)", 24), ("CS negative day-pairs", 21)]
    for index, (name, width) in enumerate(head, start=1):
        ladder.cell(3, index, name).font = Font(bold=True)
        ladder.column_dimensions[get_column_letter(index)].width = width
    for offset, row_data in enumerate(table.itertuples()):
        row = 4 + offset
        ladder.cell(row, 1, row_data.symbol)
        ladder.cell(row, 2, round(row_data.adv_cr, 3)).number_format = "0.000"
        ladder.cell(row, 3, round(row_data.ladder_half_bps, 1)).number_format = "0.0"
        ladder.cell(row, 4, round(row_data.cs_half_bps, 1)).number_format = "0.0"
        ladder.cell(row, 5, round(row_data.cs_negative_frac, 3)).number_format = "0.0%"
    for cell in ladder[3]:
        cell.fill = PatternFill("solid", fgColor="BDD7EE")

    target = report.save(book, "Slippage.xlsx")
    print(f"\n  written: {target}")


if __name__ == "__main__":
    main()
