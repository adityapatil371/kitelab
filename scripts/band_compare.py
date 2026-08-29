"""Does the 2% band earn its keep? Same three stacks, band ON vs band OFF.

    python -m scripts.band_compare

The band is the dead zone around the EMAs: with band=2% you must close 2%
ABOVE all three EMAs to buy and 2% BELOW one of them to be forced out. With
band=0% the rule is the bare "close above/below the EMA", so price hugging a
line flips the signal constantly.

Writes output/EMA Band Comparison.xlsx:
    Summary          every stack x band: trades, win %, expectancy, PF,
                     gross, charges, net, median hold, median stop distance
    Q-M-W / M-W-D / W-D-H
                     every trade of BOTH bands in the flat format, with live
                     formulas (Shares / Cost / Profit / Cum Profit / Points)

Five assigned stocks, class close-only convention, common window per stock
(hourly data starts 2015), fixed Rs1,00,000 capital and 1% risk. GROSS is
what the trade sheets show; the Summary also carries charges and net,
because avoiding charges is the band's whole justification.
"""
from __future__ import annotations

import math
import re
import shutil
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from kitelab import report
from scripts.tf_compare import ASSIGNED, simulate_variant, window_start

STACKS = [("QMW", "Q-M-W"), ("MWD", "M-W-D"), ("WDH", "W-D-H")]
BANDS = [("2% band", 0.02), ("no band", 0.0)]
HEADERS = ["Band", "Stock", "Entry Date", "Exit Date", "Entry Price", "Stop Loss",
           "Exit", "Exit Reason", "No. of Shares", "Cost of Entry", "Profit",
           "Cum Profit", "Points"]
WIDTHS = [10, 11, 17, 17, 11, 11, 11, 16, 12, 13, 11, 12, 9]


def collect(variant: str, band: float) -> list[dict]:
    out = []
    for symbol in ASSIGNED:
        start = window_start(symbol)
        out += [t for t in simulate_variant(symbol, variant, band=band)
                if pd.Timestamp(t["entry_ts"]).normalize() >= start]
    return out


def summarise(trades: list[dict]) -> dict:
    gross = [t["gross_profit"] for t in trades]
    wins = [g for g in gross if g > 0]
    losses = [g for g in gross if g <= 0]
    charges = sum(t["charges"] for t in trades)
    return {"trades": len(trades),
            "win_rate": len(wins) / len(gross) if gross else 0.0,
            "avg_win": float(np.mean(wins)) if wins else 0.0,
            "avg_loss": float(np.mean(losses)) if losses else 0.0,
            "expectancy": sum(gross) / len(gross) if gross else 0.0,
            "pf": (sum(wins) / -sum(losses)) if losses and sum(losses) < 0 else float("nan"),
            "gross": sum(gross), "charges": charges, "net": sum(gross) - charges,
            "charge_share": 100 * charges / sum(gross) if sum(gross) else float("nan"),
            "median_days": float(np.median([t["days_held"] for t in trades])) if trades else 0.0,
            "median_stop": float(np.median([t["stop_pct"] for t in trades])) if trades else 0.0}


def main() -> None:
    data = {(v, b_label): collect(v, b) for v, _ in STACKS for b_label, b in BANDS}
    for key, trades in data.items():
        s = summarise(trades)
        print(f"  {key[0]:<4} {key[1]:<8} {s['trades']:>4} trades  win {s['win_rate']:>5.1%}  "
              f"PF {s['pf']:>5.2f}  gross {s['gross']:>10,.0f}  charges {s['charges']:>9,.0f}  "
              f"net {s['net']:>10,.0f}  median hold {s['median_days']:>4.0f}d", flush=True)

    book = Workbook()
    book.remove(book.active)
    cached: dict[tuple[str, str], float] = {}

    # ---- summary -----------------------------------------------------------
    sheet = book.create_sheet("Summary")
    sheet["A1"] = ("Does the 2% band earn its keep? Same rule, same stocks, same window -- "
                   "only the dead zone around the EMAs changes. GROSS is before charges; "
                   "NET is after Zerodha delivery charges, which is where a band pays for itself.")
    cols = ["Stack", "Band", "Trades", "Win %", "Avg Win", "Avg Loss", "Expectancy",
            "Profit Factor", "Gross", "Charges", "Net", "Charges % of Gross",
            "Median Hold (days)", "Median Stop %"]
    fmts = ["@", "@", "0", "0.0%", "#,##0", "#,##0", "#,##0", "0.00", "#,##0", "#,##0",
            "#,##0", "0.0", "0", "0.00"]
    for c, (name, width) in enumerate(zip(cols, [9, 10, 8, 8, 10, 10, 11, 8, 12, 10, 12, 11, 12, 11]), start=1):
        cell = sheet.cell(3, c, name)
        cell.font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(c)].width = width
    row = 4
    for variant, label in STACKS:
        for b_label, _ in BANDS:
            s = summarise(data[(variant, b_label)])
            vals = [label, b_label, s["trades"], s["win_rate"], round(s["avg_win"]),
                    round(s["avg_loss"]), round(s["expectancy"]), round(s["pf"], 2),
                    round(s["gross"]), round(s["charges"]), round(s["net"]),
                    round(s["charge_share"], 1), round(s["median_days"]),
                    round(s["median_stop"], 2)]
            for c, (v, fmt) in enumerate(zip(vals, fmts), start=1):
                cell = sheet.cell(row, c, v)
                cell.number_format = fmt
                if b_label == "2% band":
                    cell.font = Font(bold=True)
            row += 1
        row += 1

    # ---- one sheet per stack, both bands ----------------------------------
    for variant, label in STACKS:
        ws = book.create_sheet(label)
        ws["A1"] = "Capital"; ws["B1"] = 100000
        ws["C1"] = "Risk %";  ws["D1"] = 0.01
        for a in ("A1", "C1"):
            ws[a].font = Font(bold=True)
        for c, (name, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
            ws.cell(3, c, name).font = Font(bold=True)
            ws.column_dimensions[get_column_letter(c)].width = width
        date_fmt = "yyyy-mm-dd hh:mm" if variant == "WDH" else "yyyy-mm-dd"
        row = 4
        for b_label, _ in BANDS:
            trades = data[(variant, b_label)]
            for symbol in ASSIGNED:
                mine = sorted([t for t in trades if t["symbol"] == symbol],
                              key=lambda t: t["entry_ts"], reverse=True)
                first = True
                for t in mine:
                    d = round(t["entry_price"], 2)
                    e = round(t["stop"], 2)
                    f = round(t["exit_price"], 2)
                    ws.cell(row, 1, b_label)
                    ws.cell(row, 2, symbol)
                    ws.cell(row, 3, pd.Timestamp(t["entry_ts"]).to_pydatetime()).number_format = date_fmt
                    ws.cell(row, 4, pd.Timestamp(t["exit_ts"]).to_pydatetime()).number_format = date_fmt
                    for col, v in ((5, d), (6, e), (7, f)):
                        ws.cell(row, col, v).number_format = "0.00"
                    ws.cell(row, 8, t["exit_reason"])
                    ws.cell(row, 9, f"=MIN(ROUNDDOWN($B$1*$D$1/(E{row}-F{row}),0),ROUNDDOWN($B$1/E{row},0))").number_format = "0"
                    ws.cell(row, 10, f"=E{row}*I{row}").number_format = "#,##0.00"
                    ws.cell(row, 11, f"=(G{row}-E{row})*I{row}").number_format = "#,##0.00"
                    ws.cell(row, 12, f"=K{row}" if first else f"=L{row-1}+K{row}").number_format = "#,##0.00"
                    ws.cell(row, 13, f"=G{row}-E{row}").number_format = "0.00"
                    shares = min(math.floor(100000 * 0.01 / (d - e)), math.floor(100000 / d))
                    profit = (f - d) * shares
                    cum = profit if first else cached[(label, f"L{row-1}")] + profit
                    cached[(label, f"I{row}")] = shares
                    cached[(label, f"J{row}")] = round(d * shares, 2)
                    cached[(label, f"K{row}")] = round(profit, 2)
                    cached[(label, f"L{row}")] = round(cum, 2)
                    cached[(label, f"M{row}")] = round(f - d, 2)
                    first = False
                    row += 1

    target = report.save(book, "EMA Band Comparison.xlsx")

    # ---- fill the cached values so formulas DISPLAY in every viewer --------
    names = {f"xl/worksheets/sheet{i}.xml": n
             for i, n in enumerate(book.sheetnames, start=1)}
    tmp = str(target) + ".tmp"
    with zipfile.ZipFile(target) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            blob = zin.read(item)
            if item in names:
                name = names[item]
                xml = blob.decode()

                def fill(m, name=name):
                    v = cached.get((name, m.group(1)))
                    return m.group(0) if v is None else m.group(0).replace("<v />", f"<v>{v}</v>")

                xml, n = re.subn(r'<c r="([A-Z]+\d+)"[^>]*><f>[^<]*</f><v /></c>', fill, xml)
                if n:
                    print(f"    {name}: {n} formula cells given cached values")
                blob = xml.encode()
            zout.writestr(item, blob)
    shutil.move(tmp, target)
    print(f"\n  written: {target}")


if __name__ == "__main__":
    main()
