"""Darvas 20/10 -- in on a 20-candle high, out on a 10-candle low.

    python -m scripts.darvas_test

Four questions, in order:
  1. how does the rule as specified compare with the strategies already in the lab
  2. does it hold up on the 150 stocks nothing was ever tuned on
  3. how much of the answer depends on the two window lengths
  4. does it survive realistic fills, or is the edge inside the spread

Both halves are reported throughout. The trade table sums every signal as if
separately funded at a fixed Rs1,00,000; the account table is one Rs2,50,000 pot
at 1% risk that can only take what it can afford. They answer different
questions and they do not always agree.
"""
from __future__ import annotations

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from kitelab import backtest, config, darvas, portfolio, report, slippage, strategies

CAPITAL = 250_000.0
RISK = 0.01
WINDOWS = [(10, 5), (20, 10), (20, 20), (40, 20), (55, 20)]


def collect(build, symbols) -> list[dict]:
    out = []
    for symbol in symbols:
        try:
            out.extend(build(symbol))
        except (SystemExit, FileNotFoundError):
            pass
    return out


def summarise(trades: list[dict], subset=None) -> dict:
    if subset is not None:
        trades = [t for t in trades if t["symbol"] in subset]
    if not trades:
        return {"trades": 0}
    net = np.array([t["net_profit"] for t in trades], dtype=float)
    wins, losses = net[net > 0], net[net <= 0]
    gross_win, gross_loss = float(wins.sum()), float(-losses.sum())
    account = portfolio.run(trades, CAPITAL, RISK)
    return {
        "trades": len(trades),
        "win_rate": round(100 * len(wins) / len(net), 1),
        "avg_win": round(float(wins.mean()) if len(wins) else 0.0),
        "avg_loss": round(float(losses.mean()) if len(losses) else 0.0),
        "expectancy": round(float(net.mean())),
        "net": round(float(net.sum())),
        "pf": round(gross_win / gross_loss, 2) if gross_loss else None,
        "median_bars": int(np.median([t["bars_held"] for t in trades])),
        "cagr": round(account["cagr_pct"], 2),
        "maxdd": round(account["max_drawdown_pct"], 1),
        "taken": len(account["taken"]),
    }


def line(label, s):
    if not s.get("trades"):
        return f"  {label:<34}  no trades"
    return (f"  {label:<34}{s['trades']:>7}{s['win_rate']:>7}{s['expectancy']:>11,}"
            f"{s['net']:>13,}{str(s['pf']):>6}{s['median_bars']:>7}"
            f"{s['cagr']:>8.1f}%{s['maxdd']:>8.1f}%")


HEAD = (f"  {'':<34}{'trades':>7}{'win%':>7}{'expect':>11}{'net':>13}"
        f"{'PF':>6}{'bars':>7}{'CAGR':>9}{'maxDD':>9}")


def main() -> None:
    cfg = config.load()
    symbols = cfg.all_symbols
    rows = []

    print("\n  1. the rule as specified, against what is already in the lab")
    print(HEAD)
    print("  " + "-" * 99)
    darvas_all = collect(lambda s: darvas.simulate(s), symbols)
    peers = [
        ("Darvas 20/10 (daily closes)", darvas_all),
        ("EMA stack M/W/D, 2% band", collect(lambda s: backtest.simulate(s), symbols)),
        ("ATH Breakout, swing-low trail",
         collect(lambda s: strategies.ath_breakout_trades(s, trailing_stops=True), symbols)),
    ]
    for label, trades in peers:
        stats = summarise(trades)
        rows.append(("Against the other strategies", label, stats))
        print(line(label, stats), flush=True)

    print("\n  2. Darvas 20/10 by universe -- 'holdout' is the 150 nothing was tuned on")
    print(HEAD)
    print("  " + "-" * 99)
    universes = [("All 199 stocks", None), ("49 in-sample", set(cfg.in_sample)),
                 ("150 holdout", set(cfg.out_of_sample))]
    for label, members in universes:
        stats = summarise(darvas_all, members)
        rows.append(("By universe", label, stats))
        print(line(label, stats), flush=True)

    print("\n  3. how much rests on the two window lengths")
    print(HEAD)
    print("  " + "-" * 99)
    for entry_len, exit_len in WINDOWS:
        trades = (darvas_all if (entry_len, exit_len) == (20, 10)
                  else collect(lambda s, a=entry_len, b=exit_len:
                               darvas.simulate(s, a, b), symbols))
        label = f"in {entry_len} / out {exit_len}" + ("  <- as specified"
                                                      if (entry_len, exit_len) == (20, 10) else "")
        stats = summarise(trades)
        rows.append(("Window lengths", label, stats))
        print(line(label, stats), flush=True)

    print("\n  4. does the edge survive being executed")
    print(HEAD)
    print("  " + "-" * 99)
    stats = summarise(darvas_all)
    rows.append(("Execution", "Perfect fills", stats))
    print(line("Perfect fills", stats), flush=True)
    slippage.ENABLED, slippage.MAX_PARTICIPATION = True, 0.01
    slippage.reset()
    real = collect(lambda s: darvas.simulate(s), symbols)
    stats = summarise(real)
    rows.append(("Execution", "Realistic fills (spread + impact + 1% size cap)", stats))
    print(line("Realistic fills, 1% size cap", stats), flush=True)
    intrabar_label = "Turtle reading: channel edges as live orders"
    slippage.ENABLED, slippage.MAX_PARTICIPATION = False, None
    slippage.reset()
    stats = summarise(collect(lambda s: darvas.simulate(s, intrabar=True), symbols))
    rows.append(("Execution", intrabar_label, stats))
    print(line("Intrabar (Turtle reading)", stats), flush=True)

    book = Workbook()
    book.remove(book.active)
    sheet = book.create_sheet("Darvas 20-10")
    sheet["A1"] = (
        "Darvas as specified: enter when the close beats the highest high of the previous "
        "20 candles, leave when it closes at or below the lowest low of the previous 10. "
        "Both windows exclude the current candle. Stop at entry is the 10-candle low, and "
        "because that line is recomputed daily it ratchets up behind a rising stock -- the "
        "exit IS the trailing stop. Trade columns sum every signal at a fixed Rs1,00,000; "
        "account columns are one Rs2,50,000 pot at 1% risk. Perfect fills unless the row "
        "says otherwise.")
    columns = [("Section", 28), ("Variant", 44), ("Trades", 9), ("Win %", 9),
               ("Avg win", 11), ("Avg loss", 11), ("Expectancy", 12), ("Total net", 14),
               ("Profit factor", 13), ("Median bars", 12), ("Account CAGR %", 15),
               ("Account max DD %", 17), ("Trades taken", 13)]
    for index, (name, width) in enumerate(columns, start=1):
        sheet.cell(3, index, name).font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    for offset, (section, label, s) in enumerate(rows):
        row = 4 + offset
        values = [section, label, s.get("trades"), s.get("win_rate"), s.get("avg_win"),
                  s.get("avg_loss"), s.get("expectancy"), s.get("net"), s.get("pf"),
                  s.get("median_bars"), s.get("cagr"), s.get("maxdd"), s.get("taken")]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row, index, value)
            if index in (5, 6, 7, 8):
                cell.number_format = "#,##0"
            if index in (4, 11, 12):
                cell.number_format = "0.0"
    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor="BDD7EE")

    print(f"\n  written: {report.save(book, 'Darvas 20-10.xlsx')}")


if __name__ == "__main__":
    main()
