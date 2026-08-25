"""Backtest the two level-based strategies and write them to Excel.

    python -m scripts.level_trades

One sheet per strategy with all stocks in it, matching the layout of your own file.
Trades run newest-first within each stock; cumulative columns run down the sheet in
display order, the way yours do.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kitelab import config, strategies

OUTPUT = Path(__file__).resolve().parent.parent / "output"
HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
LOSS_FONT = Font(color="C00000")

RULES = {
    "Support Bounce": (
        "RULE : 1. price touches a support line. 2. if the candle is red, check the next "
        "one. 3. the first green candle whose body midpoint is above the line is the "
        "signal. 4. buy-stop at that candle's high, stop-loss at its low, target 1.5R. "
        "5. abandon if price closes decisively below the line while waiting. "
        "30-minute candles. No signal before the level's second touch (valid_from)."
    ),
    "Breakout": (
        "RULE : 1. a resistance level already touched twice is confirmed. 2. the first "
        "30-minute close above it is the break. 3. buy-stop at that candle's high, "
        "target 1.5R. 4. stop-loss is the last DAILY pivot low already confirmed at "
        "entry time (a 5-bar pivot is not visible until 5 bars later). "
        "No signal before the level's second touch (valid_from)."
    ),
}

COLUMNS = [
    ("Trade #", "trade_no", "0", 8), ("Date", "entry_date", "yyyy-mm-dd", 11),
    ("Time", "entry_time", "@", 7), ("Stock Name", "symbol", "@", 11),
    ("Level", "level", "#,##0.00", 10), ("Max Risk Capital", "max_risk_capital", "#,##0", 11),
    ("Entry Price", "entry_price", "#,##0.00", 11), ("Stop Loss Price", "stop", "#,##0.00", 12),
    ("Range", "range", "#,##0.00", 9), ("Desired R:R", "reward_ratio", "0.0", 10),
    ("Exit", "target", "#,##0.00", 10), ("Risk Per Share", "risk_per_share", "#,##0.00", 12),
    ("No. of Shares", "shares", "#,##0", 12), ("Cost of Entry", "cost_of_entry", "#,##0.00", 13),
    ("Cum Cost", "cum_cost", "#,##0.00", 13), ("Actual Exit", "exit_price", "#,##0.00", 11),
    ("Exit Date", "exit_date", "yyyy-mm-dd", 11), ("Exit Reason", "exit_reason", "@", 18),
    ("Profit", "gross_profit", "#,##0.00", 11), ("Cum Profit", "cum_gross", "#,##0.00", 12),
    ("R Multiple", "r_multiple", "0.00", 10), ("Charges", "charges", "#,##0.00", 10),
    ("Net Profit", "net_profit", "#,##0.00", 11), ("Cum Net Profit", "cum_net", "#,##0.00", 13),
]


def _order_and_accumulate(trades: list[dict], symbols: list[str]) -> list[dict]:
    ordered: list[dict] = []
    for symbol in symbols:
        mine = [t for t in trades if t["symbol"] == symbol]
        ordered.extend(sorted(mine, key=lambda t: t["entry_ts"], reverse=True))
    cost = gross = net = 0.0
    for number, trade in enumerate(ordered, start=1):
        cost += trade["cost_of_entry"]
        gross += trade["gross_profit"]
        net += trade["net_profit"]
        trade.update(trade_no=number, cum_cost=cost, cum_gross=gross, cum_net=net)
    return ordered


def _write_sheet(book: Workbook, title: str, rule: str, trades: list[dict]) -> None:
    sheet = book.create_sheet(title)
    sheet["A1"] = rule
    sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=len(COLUMNS))
    sheet.row_dimensions[1].height = 34

    for index, (heading, _, _, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=4, column=index, value=heading)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width

    for offset, trade in enumerate(trades):
        for index, (_, key, number_format, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=5 + offset, column=index, value=trade.get(key))
            cell.number_format = number_format
            value = trade.get(key)
            if key in ("gross_profit", "net_profit", "r_multiple") and isinstance(
                    value, (int, float)) and value < 0:
                cell.font = LOSS_FONT
    sheet.freeze_panes = "A5"
    if not trades:
        sheet["A5"] = "No trades. Either no level qualified, or none produced a signal."


def _summary(name: str, trades: list[dict]) -> dict:
    if not trades:
        return {"strategy": name, "trades": 0}
    net = [t["net_profit"] for t in trades]
    wins = [v for v in net if v > 0]
    reasons: dict[str, int] = {}
    for trade in trades:
        reasons[trade["exit_reason"]] = reasons.get(trade["exit_reason"], 0) + 1
    return {
        "strategy": name, "trades": len(trades), "wins": len(wins),
        "win_rate_pct": 100 * len(wins) / len(trades),
        "gross_profit": sum(t["gross_profit"] for t in trades),
        "charges": sum(t["charges"] for t in trades),
        "net_profit": sum(net),
        "expectancy": sum(net) / len(trades),
        "avg_r": sum(t["r_multiple"] for t in trades) / len(trades),
        "peak_capital": max(t["cost_of_entry"] for t in trades),
        "exits": ", ".join(f"{k}={v}" for k, v in sorted(reasons.items())),
    }


def _write_summary(book: Workbook, rows: list[dict]) -> None:
    sheet = book.create_sheet("Summary", 0)
    fields = [("Strategy", "strategy", "@", 18), ("Trades", "trades", "0", 9),
              ("Wins", "wins", "0", 8), ("Win Rate %", "win_rate_pct", "0.0", 11),
              ("Gross Profit", "gross_profit", "#,##0.00", 13),
              ("Charges", "charges", "#,##0.00", 11),
              ("Net Profit", "net_profit", "#,##0.00", 12),
              ("Expectancy / Trade", "expectancy", "#,##0.00", 16),
              ("Avg R", "avg_r", "0.00", 9),
              ("Peak Capital", "peak_capital", "#,##0.00", 13),
              ("Exit Breakdown", "exits", "@", 46)]
    for index, (heading, _, _, width) in enumerate(fields, start=1):
        cell = sheet.cell(row=1, column=index, value=heading)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    for offset, row in enumerate(rows):
        for index, (_, key, number_format, _) in enumerate(fields, start=1):
            cell = sheet.cell(row=2 + offset, column=index, value=row.get(key))
            cell.number_format = number_format
            value = row.get(key)
            if key in ("net_profit", "expectancy") and isinstance(value, (int, float)) \
                    and value < 0:
                cell.font = LOSS_FONT
    sheet.freeze_panes = "A2"


def main() -> None:
    argparse.ArgumentParser(description=__doc__).parse_args()
    cfg = config.load()
    OUTPUT.mkdir(exist_ok=True)
    book = Workbook()
    book.remove(book.active)

    builders = {
        "Support Bounce": strategies.support_bounce_trades,
        "Breakout": strategies.breakout_trades,
    }
    summaries = []
    print()
    for name, build in builders.items():
        raw: list[dict] = []
        for symbol in cfg.symbols:
            found = build(symbol)
            raw.extend(found)
            print(f"  {name:<16} {symbol:<10} {len(found):>3} raw signals")
        kept = strategies.drop_overlaps(raw)
        dropped = len(raw) - len(kept)
        trades = _order_and_accumulate(kept, cfg.symbols)
        _write_sheet(book, name, RULES[name], trades)
        summaries.append(_summary(name, trades))
        print(f"  {name:<16} {'TOTAL':<10} {len(kept):>3} trades "
              f"({dropped} dropped as overlapping)\n")

    _write_summary(book, summaries)
    target = OUTPUT / "Level Strategies Backtest.xlsx"
    book.save(target)
    print(f"  written: {target}\n")


if __name__ == "__main__":
    main()
