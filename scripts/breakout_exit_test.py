"""Does the ATH breakout do better if you leave on a close below the 20-EMA?

    python -m scripts.breakout_exit_test

The breakout's original exit is a ratcheting daily swing-low stop: the trade ends
only when price takes out a confirmed higher low. This asks whether a simpler,
faster rule -- out on the first daily close under the 20-EMA -- keeps more of the
move or less. Three exits, same entries, same pivot-low disaster stop underneath:

    trail          ratcheting daily swing lows (the rule as taught)
    ema20          out on the first daily close below the 20-EMA, stop never moves
    trail+ema20    both live, whichever fires first

Read both halves. The trade table sums every signal as if each were separately
funded; the account table is one Rs2,50,000 pot that can only hold what it can
afford. They answer different questions and they do not always agree.
"""
from __future__ import annotations

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from kitelab import config, portfolio, report, strategies

CAPITAL = 250_000.0
RISK = 0.01
RULES = [("trail", "Swing-low trail (original)"),
         ("ema20", "Close below 20-EMA"),
         ("trail+ema20", "Trail AND 20-EMA, first to fire")]


def collect(rule: str, symbols) -> list[dict]:
    trades = []
    for symbol in symbols:
        try:
            trades.extend(strategies.ath_breakout_trades(symbol, trailing_stops=True,
                                                         exit_rule=rule))
        except (SystemExit, FileNotFoundError):
            pass
    return trades


def summarise(trades: list[dict]) -> dict:
    if not trades:
        return {}
    net = np.array([t["net_profit"] for t in trades])
    wins, losses = net[net > 0], net[net <= 0]
    held = np.array([t["bars_held"] for t in trades])
    reasons: dict[str, int] = {}
    for t in trades:
        reasons[t["exit_reason"]] = reasons.get(t["exit_reason"], 0) + 1
    gross_win = float(wins.sum())
    gross_loss = float(-losses.sum())
    return {
        "trades": len(trades),
        "win_rate": round(100 * len(wins) / len(net), 1),
        "avg_win": round(float(wins.mean()) if len(wins) else 0.0),
        "avg_loss": round(float(losses.mean()) if len(losses) else 0.0),
        "expectancy": round(float(net.mean())),
        "net": round(float(net.sum())),
        "pf": round(gross_win / gross_loss, 2) if gross_loss else None,
        "avg_r": round(float(np.mean([t["r_multiple"] for t in trades])), 2),
        "median_bars": int(np.median(held)),
        "reasons": reasons,
    }


def main() -> None:
    symbols = config.load().all_symbols
    rows = []
    print(f"\n  {'exit rule':<34}{'trades':>8}{'win%':>7}{'expectancy':>12}"
          f"{'net':>14}{'PF':>6}{'bars':>7}")
    print("  " + "-" * 88)
    for rule, label in RULES:
        trades = collect(rule, symbols)
        stats = summarise(trades)
        account = portfolio.run(trades, CAPITAL, RISK)
        rows.append((rule, label, stats, account, trades))
        print(f"  {label:<34}{stats['trades']:>8}{stats['win_rate']:>7}"
              f"{stats['expectancy']:>12,}{stats['net']:>14,}"
              f"{stats['pf']:>6}{stats['median_bars']:>7}", flush=True)

    print(f"\n  {'exit rule':<34}{'CAGR':>8}{'maxDD':>9}{'taken':>8}{'final':>14}")
    print("  " + "-" * 75)
    for _rule, label, _stats, account, _trades in rows:
        cagr = ("  wiped" if account["cagr_pct"] is None
                else f"{account['cagr_pct']:>6.1f}%")
        print(f"  {label:<34}{cagr:>8}{account['max_drawdown_pct']:>8.1f}%"
              f"{len(account['taken']):>8}{account['final']:>14,.0f}")

    print("\n  how each rule ended its trades:")
    for _rule, label, stats, _account, _trades in rows:
        ordered = sorted(stats["reasons"].items(), key=lambda kv: -kv[1])
        print(f"    {label}: " + ", ".join(f"{k} {v}" for k, v in ordered))

    book = Workbook()
    book.remove(book.active)
    sheet = book.create_sheet("Exit rules")
    sheet["A1"] = (
        "ATH breakout, all 199 stocks, same entries and the same pivot-low stop in "
        "every column -- only the way the trade ENDS changes. Trade columns sum every "
        "signal as if separately funded; account columns are one Rs2,50,000 pot at 1% "
        "risk that can only take what it can afford. Perfect fills (no spread, no "
        "impact), so read the gaps between columns, not the absolute levels.")
    columns = [("Exit rule", 34), ("Trades", 9), ("Win %", 9), ("Avg win", 11),
               ("Avg loss", 11), ("Expectancy", 12), ("Total net", 14),
               ("Profit factor", 13), ("Avg R", 9), ("Median bars held", 17),
               ("Account CAGR %", 15), ("Account max DD %", 17), ("Trades taken", 13),
               ("Final equity", 14)]
    for index, (name, width) in enumerate(columns, start=1):
        sheet.cell(3, index, name).font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    for offset, (_rule, label, stats, account, _trades) in enumerate(rows):
        row = 4 + offset
        values = [label, stats["trades"], stats["win_rate"], stats["avg_win"],
                  stats["avg_loss"], stats["expectancy"], stats["net"], stats["pf"],
                  stats["avg_r"], stats["median_bars"],
                  report.cagr_cell(account), round(account["max_drawdown_pct"], 1),
                  len(account["taken"]), round(account["final"])]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row, index, value)
            if index in (4, 5, 6, 7, 14):
                cell.number_format = "#,##0"
            if index in (3, 11, 12):
                cell.number_format = "0.0"
    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor="BDD7EE")

    detail = book.create_sheet("How trades ended")
    detail["A1"] = "Count of exit reasons under each rule."
    detail.column_dimensions["A"].width = 34
    detail.column_dimensions["B"].width = 26
    detail.column_dimensions["C"].width = 10
    for index, name in enumerate(("Exit rule", "Ended by", "Trades"), start=1):
        detail.cell(3, index, name).font = Font(bold=True)
    line = 4
    for _rule, label, stats, _account, _trades in rows:
        for reason, count in sorted(stats["reasons"].items(), key=lambda kv: -kv[1]):
            detail.cell(line, 1, label)
            detail.cell(line, 2, reason)
            detail.cell(line, 3, count)
            line += 1

    print(f"\n  written: {report.save(book, 'Breakout Exit Rules.xlsx')}")


if __name__ == "__main__":
    main()
