"""How does each strategy fare on small caps?

    python -m scripts.smallcap_test

Kite serves no fundamentals, so there is no market cap in this project and there
never will be. The stand-in is median daily TRADED VALUE over each stock's whole
history -- a liquidity measure, not a size one. They correlate strongly and the
label is only used for grouping, never inside a simulation, but do not read
"micro" as "small market cap"; read it as "hardly trades".

Two traps this script exists to avoid.

  1. Breadth. The buckets hold different numbers of stocks (52/55/50/28) and we
     already know breadth alone is worth several CAGR points, so comparing whole
     buckets measures stock COUNT as much as stock SIZE. Section 2 therefore
     draws equal-sized random baskets from each bucket.
  2. Perfect fills. The thin buckets are exactly where spread and market impact
     do their damage -- the ladder charges a Rs0.14 crore/day stock about 200bps
     a side against 2-3bps for a Rs100 crore name. Judging small caps on perfect
     fills flatters them enormously. Section 3 repeats section 2 with the real
     cost model and a 1% participation cap.

Read section 3. Sections 1 and 2 are there to show what the answer would have
been if you stopped early.
"""
from __future__ import annotations

import random

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from kitelab import config, frames, portfolio, report, signals, slippage

CAPITAL = 250_000.0
RISK = 0.01
BASKET = 25          # stocks per draw, so every bucket is judged at equal breadth
DRAWS = 30
SEED = 20260831
BUCKETS = [("Micro", 0.0, 0.5), ("Small", 0.5, 5.0), ("Mid", 5.0, 50.0),
           ("Large", 50.0, float("inf"))]
STRATEGIES = [("EMA · M/W/D", "EMA_all"), ("EMA · Q/M/W", "QMW_all"),
              ("EMA · W/D/H", "WDH_all"), ("ATH Breakout", "Breakout_all"),
              ("Darvas 20/10", "Darvas_20_10_all")]


def turnover_buckets(symbols) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {name: [] for name, _lo, _hi in BUCKETS}
    for symbol in symbols:
        try:
            day = frames.daily(symbol)
        except SystemExit:
            continue
        if len(day) < 60:
            continue
        adv = float(np.nanmedian((day["close"] * day["volume"]).to_numpy(float))) / 1e7
        for name, low, high in BUCKETS:
            if low <= adv < high:
                out[name].append(symbol)
                break
    return out


def load(cache: str, symbols) -> dict[str, list]:
    """Cached trades grouped by symbol. signals.require refuses a cache built
    from a different universe, different price files or different code."""
    by: dict[str, list] = {}
    for t in signals.require(cache, symbols):
        by.setdefault(t["symbol"], []).append(t)
    return by


def account(trades, realistic: bool) -> dict | None:
    if not trades:
        return None
    if realistic:
        slippage.ENABLED, slippage.MAX_PARTICIPATION = True, 0.01
        trades = [slippage.apply_spread(t) for t in trades]
    result = portfolio.run(trades, CAPITAL, RISK)
    slippage.ENABLED, slippage.MAX_PARTICIPATION = False, None
    net = np.array([t["net_profit"] for t in trades], dtype=float)
    wins = net[net > 0]
    losses = net[net <= 0]
    # None when the account was wiped out: it has no CAGR. Rendered as the word
    # "wiped" at the display boundary, never sorted or averaged as a 0.
    return {"cagr": result["cagr_pct"], "maxdd": result["max_drawdown_pct"],
            "trades": len(trades), "taken": len(result["taken"]),
            "win_rate": 100 * len(wins) / len(net),
            "pf": (float(wins.sum()) / -float(losses.sum())) if losses.sum() < 0 else None,
            "expectancy": float(net.mean())}


def whole_bucket(by_symbol, members, realistic) -> dict | None:
    return account([t for s in members for t in by_symbol.get(s, [])], realistic)


def sampled(by_symbol, members, realistic) -> dict | None:
    if len(members) < BASKET:
        return None
    rng = random.Random(SEED)
    cagrs, dds, pfs = [], [], []
    for _ in range(DRAWS):
        basket = rng.sample(members, BASKET)
        r = account([t for s in basket for t in by_symbol.get(s, [])], realistic)
        if r:
            cagrs.append(r["cagr"])
            dds.append(r["maxdd"])
            if r["pf"]:
                pfs.append(r["pf"])
    if not cagrs:
        return None
    # A wiped draw has no CAGR, so it sorts BELOW every draw that merely lost money
    # rather than landing on 0.0 in the middle of the distribution.
    order = sorted(cagrs, key=lambda c: (c is not None, c))
    return {"cagr": order[len(order) // 2], "p10": order[len(order) // 10],
            "p90": order[9 * len(order) // 10],
            "maxdd": sorted(dds)[len(dds) // 2],
            "pf": sorted(pfs)[len(pfs) // 2] if pfs else None,
            "draws": len(cagrs),
            "wiped": sum(1 for c in cagrs if c is None)}


def table(title, note, rows, buckets, key="cagr"):
    print(f"\n  {title}\n  {note}")
    print(f"  {'strategy':<16}" + "".join(f"{name:>13}" for name, _l, _h in BUCKETS))
    print("  " + "-" * (16 + 13 * len(BUCKETS)))
    for label, cells in rows:
        line = f"  {label:<16}"
        for name, _l, _h in BUCKETS:
            c = cells.get(name)
            if not c:
                line += f"{'—':>13}"
            elif c[key] is None:
                line += f"{'wiped':>13}"
            else:
                line += f"{c[key]:>12.1f}%"
        print(line, flush=True)


def main() -> None:
    cfg = config.load()
    buckets = turnover_buckets(cfg.all_symbols)
    print("\n  buckets by median daily traded value (a LIQUIDITY proxy, not market cap):")
    for name, low, high in BUCKETS:
        span = f"< Rs{high:g}cr" if low == 0 else (
            f"> Rs{low:g}cr" if high == float("inf") else f"Rs{low:g}-{high:g}cr")
        print(f"    {name:<7} {span:<14} {len(buckets[name]):>3} stocks")

    loaded = {label: load(cache, cfg.all_symbols) for label, cache in STRATEGIES}
    sections = []

    rows = [(label, {n: whole_bucket(loaded[label], buckets[n], False)
                     for n, _l, _h in BUCKETS}) for label, _c in STRATEGIES]
    table("1. whole buckets, perfect fills",
          "unequal stock counts, so this measures breadth as much as size", rows, buckets)
    sections.append(("1. Whole bucket, perfect fills", rows))

    rows = [(label, {n: sampled(loaded[label], buckets[n], False)
                     for n, _l, _h in BUCKETS}) for label, _c in STRATEGIES]
    table(f"2. equal baskets of {BASKET}, perfect fills",
          f"median of {DRAWS} random draws per bucket -- breadth now controlled", rows, buckets)
    sections.append((f"2. {BASKET}-stock baskets, perfect fills", rows))

    rows = [(label, {n: sampled(loaded[label], buckets[n], True)
                     for n, _l, _h in BUCKETS}) for label, _c in STRATEGIES]
    table(f"3. equal baskets of {BASKET}, REALISTIC fills",
          "spread + market impact + 1% participation cap -- the honest one", rows, buckets)
    sections.append((f"3. {BASKET}-stock baskets, realistic fills", rows))

    book = Workbook()
    book.remove(book.active)
    sheet = book.create_sheet("Small caps")
    sheet["A1"] = (
        "Every strategy against stocks grouped by how much they actually trade. There is no "
        "market cap in this project -- Kite serves no fundamentals -- so the grouping is median "
        "daily traded value over each stock's whole history. Section 1 compares whole buckets "
        "and is contaminated by breadth (52/55/50/28 stocks). Section 2 fixes that with equal "
        f"{BASKET}-stock random baskets. Section 3 repeats it with spread, market impact and a "
        "1% participation cap, and is the only one worth acting on: the thin buckets are exactly "
        "where execution costs land.")
    columns = [("Section", 32), ("Strategy", 16), ("Bucket", 10), ("Stocks in bucket", 17),
               ("CAGR %", 10), ("p10", 9), ("p90", 9), ("Max DD %", 11),
               ("Profit factor", 13), ("Trades", 9)]
    for index, (name, width) in enumerate(columns, start=1):
        sheet.cell(3, index, name).font = Font(bold=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor="BDD7EE")
    line = 4
    for section, rows in sections:
        for label, cells in rows:
            for name, _l, _h in BUCKETS:
                c = cells.get(name)
                if not c:
                    continue
                pc = lambda v: (report.WIPED_LABEL if v is None else round(v, 2))
                values = [section, label, name, len(buckets[name]),
                          pc(c["cagr"]), pc(c.get("p10", c["cagr"])),
                          pc(c.get("p90", c["cagr"])), round(c["maxdd"], 1),
                          round(c["pf"], 2) if c.get("pf") else None, c.get("trades")]
                for index, value in enumerate(values, start=1):
                    cell = sheet.cell(line, index, value)
                    if index in (5, 6, 7, 8):
                        cell.number_format = "0.0"
                line += 1

    print(f"\n  written: {report.save(book, 'Small Caps.xlsx')}")


if __name__ == "__main__":
    main()
