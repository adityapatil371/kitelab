"""Generate the EMA-strategy trade log as an Excel workbook.

    python -m scripts.ema_trades
    python -m scripts.ema_trades --count 25 --ema 20 --shares 10

One sheet per stock, newest trade first, laid out like your existing sheet: the rule
text on row 1, headers on row 4, trades from row 5. A Summary sheet carries the stats
and the point-in-time verification result.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kitelab import backtest, config

OUTPUT = Path(__file__).resolve().parent.parent / "output"

RULE = (
    "RULE : 1. buy on EOD when close is above the 20-EMA on monthly, weekly AND daily. "
    "2. stop loss is the low of the entry day. "
    "3. exit at whichever comes first -- the stop being touched (filled at the open if "
    "the day gaps through it), or the EMA stack breaking (filled at that day's close). "
    "4. re-entry only on a fresh signal. "
    "Charges are Zerodha equity delivery rates: STT 0.1% both sides, NSE txn 0.00307%, "
    "SEBI Rs10/crore, GST 18%, stamp 0.015% on buy, DP Rs15.34 on sell."
)

COLUMNS = [
    ("Trade #", "trade_no", "0", 8),
    ("Entry Date", "entry_date", "yyyy-mm-dd", 12),
    ("Entry Price", "entry_price", "#,##0.00", 12),
    ("Stop Loss", "stop", "#,##0.00", 11),
    ("Risk Per Share", "risk_per_share", "#,##0.00", 14),
    ("Exit Date", "exit_date", "yyyy-mm-dd", 12),
    ("Exit Price", "exit_price", "#,##0.00", 11),
    ("Exit Reason", "exit_reason", "@", 17),
    ("Sessions Held", "sessions_held", "0", 13),
    ("No. of Shares", "shares", "0", 13),
    ("Cost of Entry", "cost_of_entry", "#,##0.00", 13),
    ("Cum Cost", "cum_cost", "#,##0.00", 13),
    ("Profit", "gross_profit", "#,##0.00", 11),
    ("Cum Profit", "cum_gross", "#,##0.00", 12),
    ("R Multiple", "r_multiple", "0.00", 11),
    ("Charges", "charges", "#,##0.00", 10),
    ("Net Profit", "net_profit", "#,##0.00", 11),
    ("Cum Net Profit", "cum_net", "#,##0.00", 14),
    ("Monthly Bars", "months_done", "0", 13),
]

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
LOSS_FONT = Font(color="C00000")


def _write_stock_sheet(book: Workbook, symbol: str, trades: list[dict]) -> None:
    sheet = book.create_sheet(symbol)
    sheet["A1"] = RULE
    sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=len(COLUMNS))
    sheet.row_dimensions[1].height = 30

    for index, (title, _, _, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=4, column=index, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width

    for offset, trade in enumerate(trades):
        row = 5 + offset
        for index, (_, key, number_format, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row, column=index, value=trade.get(key))
            cell.number_format = number_format
            if key in ("gross_profit", "net_profit", "r_multiple") and (trade.get(key) or 0) < 0:
                cell.font = LOSS_FONT

    sheet.freeze_panes = "A5"
    if not trades:
        sheet["A5"] = "No closed trades in the available history."


def _write_summary(book: Workbook, rows: list[dict], checks: list[tuple]) -> None:
    sheet = book.create_sheet("Summary", 0)
    fields = [
        ("Stock", "symbol", "@"), ("Trades", "trades", "0"),
        ("Wins", "wins", "0"), ("Losses", "losses", "0"),
        ("Win Rate %", "win_rate_pct", "0.0"),
        ("Gross Profit", "gross_profit", "#,##0.00"),
        ("Charges", "charges", "#,##0.00"),
        ("Net Profit", "net_profit", "#,##0.00"),
        ("Avg Win", "avg_win", "#,##0.00"), ("Avg Loss", "avg_loss", "#,##0.00"),
        ("Expectancy / Trade", "expectancy_per_trade", "#,##0.00"),
        ("Avg R", "avg_r", "0.00"),
        ("Best", "best", "#,##0.00"), ("Worst", "worst", "#,##0.00"),
        ("Max Drawdown", "max_drawdown", "#,##0.00"),
        ("Peak Capital", "capital_deployed", "#,##0.00"),
        ("Stopped Out", "stopped_out", "0"), ("EMA Breaks", "ema_breaks", "0"),
    ]
    for index, (title, _, _) in enumerate(fields, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = max(11, len(title) + 2)

    for offset, summary in enumerate(rows):
        for index, (_, key, number_format) in enumerate(fields, start=1):
            cell = sheet.cell(row=2 + offset, column=index, value=summary.get(key))
            cell.number_format = number_format
            if key in ("net_profit", "expectancy_per_trade", "worst", "max_drawdown") \
                    and isinstance(summary.get(key), (int, float)) and summary[key] < 0:
                cell.font = LOSS_FONT

    start = 3 + len(rows)
    sheet.cell(row=start, column=1, value="Point-in-time verification").font = Font(bold=True)
    sheet.cell(row=start + 1, column=1, value=(
        "Each row: sessions where the fast signal agreed with screener.context, which "
        "truncates the daily data and rebuilds the weekly/monthly frames from scratch. "
        "Anything below 100% means the backtest is using lookahead."
    ))
    for offset, (symbol, agree, checked) in enumerate(checks):
        row = start + 3 + offset
        share = f"{100 * agree / checked:.2f}%" if checked else "n/a"
        sheet.cell(row=row, column=1, value=symbol)
        sheet.cell(row=row, column=2, value=f"{agree}/{checked} sessions agree ({share})")
    sheet.freeze_panes = "A2"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--count", type=int, default=25, help="trades per stock")
    parser.add_argument("--ema", type=int, default=backtest.EMA_LENGTH)
    parser.add_argument("--shares", type=int, default=backtest.SHARES)
    parser.add_argument("--verify", type=int, default=200,
                        help="sessions to cross-check against screener.context")
    args = parser.parse_args()

    cfg = config.load()
    OUTPUT.mkdir(exist_ok=True)
    book = Workbook()
    book.remove(book.active)

    summaries, checks = [], []
    print()
    for symbol in cfg.symbols:
        every = backtest.simulate(symbol, args.ema, args.shares)
        trades = backtest.recent_trades(symbol, args.count, args.ema, args.shares)
        _write_stock_sheet(book, symbol, trades)
        summaries.append(backtest.summarise(symbol, trades))

        agree, checked = backtest.verify_against_screener(symbol, args.verify, args.ema)
        checks.append((symbol, agree, checked))
        flag = "" if checked and agree == checked else "   <-- MISMATCH"
        print(f"  {symbol:<10} {len(every):>4} closed trades in history, "
              f"showing {len(trades):>2}   verify {agree}/{checked}{flag}")

    _write_summary(book, summaries, checks)
    target = OUTPUT / "EMA Strategy Backtest.xlsx"
    book.save(target)
    print(f"\n  written: {target}\n")


if __name__ == "__main__":
    main()
