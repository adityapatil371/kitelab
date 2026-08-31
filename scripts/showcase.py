"""The five assigned stocks, 25 most recent EMA trades each, in the practice-file
format: one sheet per stock, your original columns, and LIVE FORMULAS instead of
pasted numbers -- so every value shows where it came from, and changing the Capital
or Risk cells at the top recalculates the whole sheet.

    python -m scripts.showcase
    python -m scripts.showcase --count 25

Only raw market facts are written as numbers: Date, Entry Price, Stop Loss, Exit.
Everything else is an Excel formula:

    No. of Shares = MIN( ROUNDDOWN(Capital*Risk% / (Entry - Stop), 0),
                         ROUNDDOWN(Capital / Entry, 0) )
    Cost of Entry = Entry * Shares
    Cum Cost      = previous Cum Cost + Cost of Entry
    Profit        = (Exit - Entry) * Shares
    Cum Profit    = previous Cum Profit + Profit

Charges are deliberately left out to match the practice file; the net-of-charges
version of these same trades lives in EMA Backtest.xlsx.
"""
from __future__ import annotations

import argparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from kitelab import backtest, report

ASSIGNED = ["HAL", "HINDZINC", "HYUNDAI", "IRFC", "INDHOTEL"]

RULE = (
    "RULE : 1. buy at the close when price is 2% above the 20-EMA on the monthly, "
    "weekly AND daily timeframes. 2. stop loss = the entry day's low. 3. sell when the "
    "stop is hit, or when the close falls 2% below any of the three EMAs, whichever "
    "comes first. Shares = risk budget / (entry - stop), capped by capital -- see the "
    "formula in every No. of Shares cell."
)

HEADERS = ["Trade #", "Date", "Entry Price", "Stop Loss", "Exit",
           "No. of Shares", "Cost of Entry", "Cum Cost", "Profit", "Cum Profit"]
WIDTHS = [8, 12, 12, 12, 12, 13, 14, 14, 12, 12]
HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
LOSS_FONT = Font(color="C00000")


def write_sheet(book: Workbook, symbol: str, trades: list[dict],
                values: report.FormulaValues) -> None:
    sheet = book.create_sheet(symbol)

    sheet["A1"] = RULE
    sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    sheet.row_dimensions[1].height = 44

    sheet["A2"] = "Capital"
    sheet["B2"] = 100000
    sheet["C2"] = "Risk %"
    sheet["D2"] = 0.01
    sheet["D2"].number_format = "0%"
    sheet["E2"] = "<- change these two cells and every formula below recalculates"
    sheet["E2"].font = Font(italic=True, color="808080")
    for anchor in ("A2", "C2"):
        sheet[anchor].font = Font(bold=True)

    for column, (title, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = sheet.cell(row=4, column=column, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.column_dimensions[chr(64 + column)].width = width

    running_cost = running_profit = 0.0
    for offset, trade in enumerate(trades):
        row = 5 + offset
        # Raw market facts -- the only hard numbers on the sheet.
        sheet.cell(row=row, column=1, value=offset + 1)
        date_cell = sheet.cell(row=row, column=2, value=trade["entry_ts"].date())
        date_cell.number_format = "yyyy-mm-dd"
        for column, key in ((3, "entry_price"), (4, "stop"), (5, "exit_price")):
            sheet.cell(row=row, column=column, value=round(trade[key], 2)).number_format = "0.00"

        # Everything derived is a formula -- and every formula also carries the value
        # Python already computed, so the column reads without Excel recalculating
        # it (report.FormulaValues). The formula itself stays live.
        shares = trade["shares"]
        cost = trade["entry_price"] * shares
        gross = trade["gross_profit"]
        running_cost += cost
        running_profit += gross
        values.write(sheet, row, 6,
                     f"=MIN(ROUNDDOWN($B$2*$D$2/(C{row}-D{row}),0),"
                     f"ROUNDDOWN($B$2/C{row},0))", shares, "0")
        values.write(sheet, row, 7, f"=C{row}*F{row}", round(cost, 2), "#,##0.00")
        cum_cost = f"=G{row}" if offset == 0 else f"=H{row - 1}+G{row}"
        values.write(sheet, row, 8, cum_cost, round(running_cost, 2), "#,##0.00")
        values.write(sheet, row, 9, f"=(E{row}-C{row})*F{row}", round(gross, 2),
                     "#,##0.00")
        cum_profit = f"=I{row}" if offset == 0 else f"=J{row - 1}+I{row}"
        values.write(sheet, row, 10, cum_profit, round(running_profit, 2), "#,##0.00")
        if trade["exit_price"] < trade["entry_price"]:
            sheet.cell(row=row, column=9).font = LOSS_FONT

    total_row = 5 + len(trades) + 1
    sheet.cell(row=total_row, column=8, value="Total").font = Font(bold=True)
    total = values.write(sheet, total_row, 9, f"=SUM(I5:I{4 + len(trades)})",
                         round(running_profit, 2), "#,##0.00")
    total.font = Font(bold=True)
    sheet.freeze_panes = "A5"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--count", type=int, default=25, help="trades per stock")
    args = parser.parse_args()

    book = Workbook()
    book.remove(book.active)
    values = report.FormulaValues()
    print()
    for symbol in ASSIGNED:
        try:
            closed = backtest.simulate(symbol)
        except SystemExit as exc:
            print(f"  {symbol:<10} {exc}")
            continue
        # Newest first, matching the practice file's layout.
        recent = sorted(closed[-args.count:], key=lambda t: t["entry_ts"], reverse=True)
        write_sheet(book, symbol, recent, values)
        gross = sum(t["gross_profit"] for t in recent)
        print(f"  {symbol:<10} {len(recent):>2} trades   gross {gross:>10,.0f}")

    target = report.save(book, "EMA Showcase.xlsx")
    values.inject(target, book)
    print(f"\n  written: {target}\n")


if __name__ == "__main__":
    main()
