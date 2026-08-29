"""Comprehensive drawdown analysis across everything we trade.

    python -m scripts.drawdown_report

Writes output/Drawdown Analysis.xlsx:

    Read Me          -- what drawdown is, how it is measured, predictions vs verdicts
    Stock Portfolios -- EMA & Breakout accounts on 199 NSE stocks: risk sweep x
                        universe (all / 49 in-sample / 150 holdout), true daily
                        mark-to-market drawdown, time spent underwater
    Worst Episodes   -- the five deepest dips of each headline account, with dates
                        and recovery times
    Assets           -- the six class instruments: strategy accounts vs buy-and-hold
    Underwater       -- charts of drawdown depth through time

Everything uses the corrected drawdown: equity re-priced daily at market closes,
each dip measured against the peak standing at that moment. Stock signal lists come
from data/signal_cache (rebuilt 2026-08-27 on full 2006+ history).
"""
from __future__ import annotations

import pickle
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

from kitelab import backtest, config, frames, portfolio, report, sizing, strategies
from scripts.full_report import explain, write_table

CACHE = Path(__file__).resolve().parent.parent / "data" / "signal_cache"
CAPITAL = 250_000
RISKS = [0.0025, 0.005, 0.01, 0.02]

ASSETS = [  # symbol, breakout timeframe, flat per-side fee (assets_report conventions)
    ("BITCOIN", "30m", 0.0010), ("NIFTY 50", "30m", 0.0005),
    ("NIFTY BANK", "30m", 0.0005), ("GOLD", "1d", 0.0005),
    ("SILVER", "1d", 0.0005), ("CRUDEOIL", "1d", 0.0005),
]

PREDICTIONS = [
    "1. The low-risk surprise (0.25% risk beating 1% on the full universe) survives on "
    "the 150 holdout stocks: better CAGR AND shallower drawdown, though holdout "
    "drawdowns run somewhat deeper than in-sample at every risk level.",
    "2. Every asset strategy-account shows a deeper true drawdown than the old at-cost "
    "method reported, but still far shallower than buy-and-hold on the same asset.",
    "3. Crude oil buy-and-hold stays the horror exhibit (-99.99%).",
]


# ------------------------------------------------------------- analytics ----

def underwater_series(curve):
    """(dates, drawdown %) below the running peak, day by day."""
    peak = float("-inf")
    dates, depth = [], []
    for day, eq in curve:
        peak = max(peak, eq)
        dates.append(day)
        depth.append(100 * (eq - peak) / peak)
    return dates, depth


def underwater_stats(curve):
    """(longest spell in days, still-running spell in days)."""
    peak = float("-inf")
    peak_day = None
    longest = current = 0
    for day, eq in curve:
        if eq >= peak:
            peak, peak_day, current = eq, day, 0
        else:
            current = (day - peak_day).days
            longest = max(longest, current)
    return longest, current


def episodes(curve, top: int = 5):
    """Each peak-to-recovery spell: depth %, dates, duration. Deepest first."""
    out = []
    peak = float("-inf")
    peak_day = trough_day = None
    trough = None
    for day, eq in curve:
        if eq >= peak:
            if trough is not None and trough < peak:
                out.append({"peak_day": peak_day, "trough_day": trough_day,
                            "depth_pct": 100 * (trough - peak) / peak,
                            "recovered": day,
                            "days": (day - peak_day).days})
            peak, peak_day = eq, day
            trough, trough_day = eq, day
        elif eq < trough:
            trough, trough_day = eq, day
    if trough is not None and trough < peak:  # still open at the end
        out.append({"peak_day": peak_day, "trough_day": trough_day,
                    "depth_pct": 100 * (trough - peak) / peak,
                    "recovered": None,
                    "days": (curve[-1][0] - peak_day).days})
    return sorted(out, key=lambda e: e["depth_pct"])[:top]


def bh_stats(daily: pd.DataFrame):
    closes = daily["close"]
    dd = closes / closes.cummax() - 1
    years = (daily["ts"].iloc[-1] - daily["ts"].iloc[0]).days / 365.25
    growth = closes.iloc[-1] / closes.iloc[0]
    curve = list(zip(daily["ts"], closes))
    longest, current = underwater_stats(curve)
    return {"cagr": 100 * (growth ** (1 / years) - 1),
            "maxdd": float(100 * dd.min()),
            "trough": daily["ts"].iloc[int(dd.values.argmin())],
            "longest_uw": longest, "current_uw": current, "years": years}


# ----------------------------------------------------------------- main ----

def main() -> None:
    cfg = config.load()
    universes = [("all 199", None),
                 ("49 in-sample", set(cfg.in_sample)),
                 ("150 holdout", set(cfg.out_of_sample))]

    ema_all = pickle.loads((CACHE / "EMA_199.pkl").read_bytes())
    brk_all = pickle.loads((CACHE / "Breakout_199.pkl").read_bytes())
    since_2015 = pd.Timestamp("2015-02-02")
    strat_windows = [
        ("EMA, 2006-2026", ema_all),
        ("EMA, 2015-2026 (Breakout's window)",
         [t for t in ema_all if pd.Timestamp(t["entry_ts"]) >= since_2015]),
        ("Breakout, 2015-2026", brk_all),
    ]

    print("\n  stock portfolios (capital 250,000):")
    stock_rows: dict[str, list[list]] = {}
    headline_curves: dict[str, list] = {}
    for label, signals in strat_windows:
        rows = []
        for uni_label, members in universes:
            subset = (signals if members is None
                      else [t for t in signals if t["symbol"] in members])
            for risk in RISKS:
                r = portfolio.run(subset, CAPITAL, risk)
                longest, current = underwater_stats(r["curve"])
                rows.append([uni_label, 100 * risk, r["cagr_pct"],
                             r["max_drawdown_pct"], r["legacy_max_drawdown_pct"],
                             longest / 365.25, current / 365.25,
                             len(r["taken"]), r["signals"], r["final"]])
                if members is None and risk in (0.0025, 0.01) \
                        and "2015" not in label.split(",")[1]:
                    headline_curves[f"{label.split(',')[0]} @ {100*risk:.2f}%"] = r["curve"]
                if members is None and risk in (0.0025, 0.01) and label.startswith("Breakout"):
                    headline_curves[f"Breakout @ {100*risk:.2f}%"] = r["curve"]
            print(f"    {label:<36} {uni_label:<14} done")
        stock_rows[label] = rows

    nifty_full = bh_stats(frames.daily("NIFTY 50"))
    nifty_2015 = bh_stats(frames.daily("NIFTY 50").query("ts >= '2015-02-02'")
                          .reset_index(drop=True))

    print("  assets:")
    asset_rows = []
    for symbol, brk_tf, fee in ASSETS:
        backtest.FLAT_FEE_RATE = fee
        sizing.FRACTIONAL = True
        try:
            ema = backtest.simulate(symbol)
            brk = strategies.ath_breakout_trades(symbol, True, timeframe=brk_tf)
            daily = frames.daily(symbol)
            bh = bh_stats(daily)
            for strat_label, trades in (("EMA", ema), ("Breakout", brk)):
                if not trades:
                    continue
                r = portfolio.run(trades, 100_000, 0.01)
                longest, current = underwater_stats(r["curve"])
                asset_rows.append([
                    symbol, strat_label, r["cagr_pct"], r["max_drawdown_pct"],
                    r["legacy_max_drawdown_pct"], longest / 365.25, current / 365.25,
                    bh["cagr"], bh["maxdd"], bh["longest_uw"] / 365.25])
        finally:
            backtest.FLAT_FEE_RATE = None
            sizing.FRACTIONAL = False
        print(f"    {symbol:<11} done")

    # ------------------------------------------------------------- book ----
    book = Workbook()
    book.remove(book.active)

    sheet = book.create_sheet("Read Me")
    sheet.column_dimensions["A"].width = 110
    row = explain(sheet, 1, "WHAT THIS FILE IS", [
        "Drawdown -- the peak-to-valley fall of the ACCOUNT, not the stock -- measured "
        "properly across everything we trade: both strategies on 199 NSE stocks, and "
        "the six class instruments (Bitcoin, two indices, three MCX futures) against "
        "buy-and-hold.",
        "Proper means: equity is re-priced EVERY trading day as cash + shares x that "
        "day's close (open positions at market, not at cost), and every dip is measured "
        "against the peak standing at that moment -- never the final peak. The old "
        "method broke all three of those rules and understated drawdowns by up to half.",
    ])
    row = explain(sheet, row, "WORKED EXAMPLE (the whole idea in four days)", [
        "You have Rs100, buy 1 share at 100. Closes: 100, 110, 70, then you sell at 80.",
        "Equity day by day: 100, 110, 70, 80. The high-water mark reached 110, the "
        "worst day sat at 70, so max drawdown = (70-110)/110 = -36%. The old at-cost "
        "method only looked at the final settlement (80 vs 100 invested = -20%) and "
        "never saw the -36% day at all. Order matters too: highest-ever minus "
        "lowest-ever is NOT a drawdown unless the peak came FIRST.",
    ])
    row = explain(sheet, row, "PREDICTIONS -- written before the numbers were computed",
                  PREDICTIONS)
    row = explain(sheet, row, "VERDICTS", [
        "1. PARTLY CORRECT. For EMA the direction survives everywhere it matters: "
        "full history all-199 (0.25% risk: 24.6% CAGR / -31.6% DD vs 1%: 12.2% / "
        "-59.9%) and on the 150 holdout stocks no parameter ever saw (27.3% / -32.6% "
        "vs 14.0% / -66.3%). It FAILS for Breakout: with only ~2,000 signals in 11 "
        "years, 0.25% risk leaves the account under-deployed (2.9% CAGR vs 5.3% at "
        "1%). The honest general rule: cutting risk per trade ALWAYS shrinks the "
        "drawdown; it only also raises returns when signals are plentiful enough to "
        "keep the freed-up cash working (EMA has ~10,000, Breakout ~2,000).",
        "2. CORRECT: every strategy account's true drawdown is deeper than the old "
        "figure, and every one is far shallower than buy-and-hold on the same asset "
        "(see Assets).",
        "3. CORRECT: crude oil buy-and-hold bottomed at ~-100% in April 2020.",
        "THE SUSTAINABILITY QUESTION: at 1% risk, drawdowns run -56% to -66% with 5-6 "
        "YEARS underwater -- the 'unsustainable' verdict stands on the pain, though "
        "under the class's close-only exit convention the returns at 1% are no longer "
        "FD-level (EMA ~12% CAGR). The refinement the data adds: risk per trade, not "
        "the strategy alone, sets the pain -- at 0.25% risk the same signals produce "
        "index-beating returns with index-like drawdowns.",
    ])
    explain(sheet, row, "CAVEATS, so nobody is fooled", [
        "Survivorship bias: the universe is TODAY'S surviving stocks, worst for the "
        "2008-2009 window -- real drawdowns and returns would be worse. No slippage.",
        "EMA signals start 2006, Breakout can only start 2015 (Kite has no older "
        "intraday data) -- the 'EMA, 2015-2026' block exists for a fair window match.",
        "MCX futures are continuous series with rollover jumps: their numbers carry "
        "extra noise. Indices are not directly tradable (futures/funds approximated).",
        "CONVENTIONS (2026-08-28): EMA uses the class rule -- the stop is checked at "
        "bar CLOSES only. Breakout keeps live intrabar stops, because its buy-stop "
        "entry is inherently an intrabar order; the two are not directly comparable "
        "on stop behaviour.",
        "The risk sweep is four values we looked at AFTER the fact -- treat '0.25% is "
        "best' as a direction with a mechanical explanation (smaller positions -> cash "
        "lasts -> nearly twice the signals taken -> smoother compounding), not a tuned "
        "constant. The holdout check makes it more believable, not proven.",
        "An earlier version of the EMA numbers (2026-08-26) accidentally used a stale "
        "signal cache that started in 2015; this file uses the full-history rebuild. "
        "With 2006-2014 included, EMA's worst episode is the 2008 crash, not 2024-26.",
    ])

    # ---- stock portfolios ----
    sheet = book.create_sheet("Stock Portfolios")
    headers = ["Universe", "Risk %", "CAGR %", "True Max DD %", "Old (at-cost) DD %",
               "Longest Underwater (yrs)", "Underwater Now (yrs)", "Trades Taken",
               "Signals", "Final Equity"]
    widths = [14, 8, 9, 13, 13, 14, 12, 12, 10, 13]
    fmts = ["@", "0.00", "0.0", "0.0", "0.0", "0.0", "0.0", "#,##0", "#,##0", "#,##0"]
    row = explain(sheet, 1, "ONE ACCOUNT, MANY STOCKS", [
        f"Starting capital Rs{CAPITAL:,}, risk per trade as shown, one position per "
        "stock, signals skipped when cash runs out. 'Underwater' = below a previous "
        "equity peak. Compare each block's 0.25% row against its 1% row -- and compare "
        "the 150-holdout block (stocks no parameter ever saw) against in-sample.",
        f"Benchmark, NIFTY 50 buy-and-hold: 2006-2026 CAGR {nifty_full['cagr']:.1f}%, "
        f"max DD {nifty_full['maxdd']:.1f}%, longest underwater "
        f"{nifty_full['longest_uw']/365.25:.1f} yrs. 2015-2026: CAGR "
        f"{nifty_2015['cagr']:.1f}%, max DD {nifty_2015['maxdd']:.1f}%. A fixed "
        "deposit: ~7% CAGR, zero drawdown -- any row below FD-level CAGR with a huge "
        "drawdown is strictly worse than doing nothing clever.",
    ], span=len(headers))
    for label, rows in stock_rows.items():
        row = write_table(sheet, row, label, headers, widths, rows, fmts) + 1

    # ---- worst episodes ----
    sheet = book.create_sheet("Worst Episodes")
    row = explain(sheet, 1, "THE FIVE DEEPEST DIPS OF EACH HEADLINE ACCOUNT", [
        "A drawdown is not one bad day -- it is a spell: the account peaks, sinks to a "
        "trough, and (maybe) climbs back. 'Recovered' empty means the account was "
        "still below that peak when the data ended. Length is peak to recovery.",
        "All accounts here: Rs250,000, all 199 stocks.",
    ], span=7)
    ep_headers = ["Peak", "Trough", "Depth %", "Recovered", "Length (yrs)"]
    ep_widths = [13, 13, 10, 13, 11]
    ep_fmts = ["@", "@", "0.0", "@", "0.0"]
    for label, curve in headline_curves.items():
        rows = [[str(e["peak_day"].date()), str(e["trough_day"].date()),
                 e["depth_pct"],
                 str(e["recovered"].date()) if e["recovered"] is not None else "not yet",
                 e["days"] / 365.25] for e in episodes(curve)]
        row = write_table(sheet, row, label, ep_headers, ep_widths, rows, ep_fmts) + 1

    # ---- assets ----
    sheet = book.create_sheet("Assets")
    a_headers = ["Asset", "Strategy", "CAGR %", "True Max DD %", "Old (at-cost) DD %",
                 "Longest Underwater (yrs)", "Underwater Now (yrs)",
                 "Buy&Hold CAGR %", "Buy&Hold Max DD %", "B&H Longest UW (yrs)"]
    a_widths = [12, 10, 9, 13, 13, 14, 12, 12, 13, 13]
    a_fmts = ["@", "@", "0.0", "0.0", "0.0", "0.0", "0.0", "0.0", "0.0", "0.0"]
    row = explain(sheet, 1, "THE SIX CLASS INSTRUMENTS", [
        "One account per asset: 100,000 currency units, 1% risk, fractional sizing, "
        "flat per-side fees (0.10% crypto, 0.05% others). Bitcoin is in US dollars. "
        "The point of the comparison: the strategies' exit rules trade away some "
        "return for MUCH shallower drawdowns than holding the asset raw.",
    ], span=len(a_headers))
    write_table(sheet, row, "Strategy accounts vs buy-and-hold",
                a_headers, a_widths, asset_rows, a_fmts)

    # ---- underwater charts ----
    sheet = book.create_sheet("Underwater")
    sheet["A1"] = ("Drawdown through time: 0 = at a new equity peak; the line hanging "
                   "below 0 is how far the account sat under its previous best. "
                   "Data columns to the right feed the charts.")
    curves = list(headline_curves.items())
    col = 8
    anchor = 3
    for label, curve in curves:
        dates, depth = underwater_series(curve)
        letter_d, letter_v = get_column_letter(col), get_column_letter(col + 1)
        sheet.cell(row=1, column=col, value="date")
        sheet.cell(row=1, column=col + 1, value=label)
        for i, (d, v) in enumerate(zip(dates, depth), start=2):
            sheet.cell(row=i, column=col, value=d).number_format = "yyyy-mm-dd"
            sheet.cell(row=i, column=col + 1, value=round(v, 2)).number_format = "0.0"
        chart = LineChart()
        chart.title = f"{label}: % below previous equity peak"
        chart.y_axis.title = "drawdown %"
        chart.height, chart.width = 9, 24
        series = Reference(sheet, min_col=col + 1, min_row=1, max_row=1 + len(dates))
        chart.add_data(series, titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=col, min_row=2,
                                       max_row=1 + len(dates)))
        sheet.add_chart(chart, f"A{anchor}")
        sheet.column_dimensions[letter_d].hidden = True
        sheet.column_dimensions[letter_v].hidden = True
        col += 2
        anchor += 19

    target = report.save(book, "Drawdown Analysis.xlsx")
    print(f"\n  written: {target}")


if __name__ == "__main__":
    main()
