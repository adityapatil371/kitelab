"""The one consolidated report -- everything in a single workbook.

    python -m scripts.master_report

Writes output/Kitelab Master Report.xlsx and replaces the pile of separate
report files (removed 2026-08-28 at Aditya's request):

    Read Me            -- conventions, caveats, how to regenerate anything
    Master Comparison  -- every account and benchmark we track, side by side
    199 Stocks         -- portfolio risk sweep x universe for both strategies
    Timeframes         -- Q/M/W vs M/W/D vs W/D/H on the five assigned stocks
    Futures & Assets   -- the six class instruments vs buy-and-hold
    Drawdowns          -- method, worked example, worst episodes, underwater charts
    BTC Validation     -- preserved copy of the manual-backtest validation (the
                          original upload no longer exists; this data is the
                          only unreproducible content in any report)

Conventions: EMA uses the CLASS rule (stops checked at bar closes only,
2026-08-28); Breakout keeps live intrabar stops (its buy-stop entry is
inherently intrabar). Stock portfolios are net of Zerodha delivery charges;
timeframe-variant tables are GROSS; assets use flat per-side fees.
"""
from __future__ import annotations

import pickle
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

from kitelab import backtest, config, frames, portfolio, report, sizing, strategies
from scripts.full_report import explain, write_table
from scripts.drawdown_report import bh_stats, episodes, underwater_series, underwater_stats
from scripts.tf_compare import (ASSIGNED, VARIANTS as TF_VARIANTS, simulate_variant,
                                summarise as tf_summarise, window_start)

CACHE = Path(__file__).resolve().parent.parent / "data" / "signal_cache"
CAPITAL = 250_000
RISKS = [0.0025, 0.005, 0.01, 0.02]
ASSETS = [("BITCOIN", "30m", 0.0010), ("NIFTY 50", "30m", 0.0005),
          ("NIFTY BANK", "30m", 0.0005), ("GOLD", "1d", 0.0005),
          ("SILVER", "1d", 0.0005), ("CRUDEOIL", "1d", 0.0005)]
BTC_VALIDATED = Path(__file__).resolve().parent.parent / "output" / "Aditya P bitcoin - validated.xlsx"


def run_stock_portfolios():
    cfg = config.load()
    universes = [("all 199", None), ("49 in-sample", set(cfg.in_sample)),
                 ("150 holdout", set(cfg.out_of_sample))]
    ema_all = pickle.loads((CACHE / "EMA_199.pkl").read_bytes())
    brk_all = pickle.loads((CACHE / "Breakout_199.pkl").read_bytes())
    since_2015 = pd.Timestamp("2015-02-02")
    blocks = [("EMA, 2006-2026 (class close-only rule)", ema_all),
              ("EMA, 2015-2026 (Breakout's window)",
               [t for t in ema_all if pd.Timestamp(t["entry_ts"]) >= since_2015]),
              ("Breakout, 2015-2026 (intrabar stops)", brk_all)]
    tables, headline = {}, {}
    for label, signals in blocks:
        rows = []
        for uni_label, members in universes:
            subset = (signals if members is None
                      else [t for t in signals if t["symbol"] in members])
            for risk in RISKS:
                r = portfolio.run(subset, CAPITAL, risk)
                longest, current = underwater_stats(r["curve"])
                rows.append([uni_label, 100 * risk, r["cagr_pct"], r["max_drawdown_pct"],
                             longest / 365.25, current / 365.25,
                             len(r["taken"]), r["final"]])
                if members is None and risk in (0.0025, 0.01):
                    headline[f"{label.split(',')[0]} @ {100*risk:.2f}%"
                             + (" (2015+)" if "2015-2026" in label and label.startswith("EMA") else "")] = r
        tables[label] = rows
        print(f"  portfolios: {label} done")
    return tables, headline


def run_timeframes():
    out = {}
    for key, label, desc in TF_VARIANTS:
        rows, allt = [], []
        for symbol in ASSIGNED:
            w = window_start(symbol)
            kept = [t for t in simulate_variant(symbol, key)
                    if pd.Timestamp(t["entry_ts"]).normalize() >= w]
            allt += kept
            s = tf_summarise(kept)
            rows.append([symbol, s["trades"], s["win_rate"], round(s["avg_win"]),
                        round(s["avg_loss"]), round(s["expectancy"]),
                        round(s["profit_factor"], 2), round(s["gross"])])
        s = tf_summarise(allt)
        rows.append(["ALL 5", s["trades"], s["win_rate"], round(s["avg_win"]),
                    round(s["avg_loss"]), round(s["expectancy"]),
                    round(s["profit_factor"], 2), round(s["gross"])])
        out[f"{label} -- {desc}"] = (rows, s)
        print(f"  timeframes: {label} done")
    return out


def run_assets():
    rows = {}
    for symbol, brk_tf, fee in ASSETS:
        backtest.FLAT_FEE_RATE = fee
        sizing.FRACTIONAL = True
        try:
            ema = backtest.simulate(symbol)
            brk = strategies.ath_breakout_trades(symbol, True, timeframe=brk_tf)
            bh = bh_stats(frames.daily(symbol))
            for strat, trades in (("EMA", ema), ("Breakout", brk)):
                if not trades:
                    continue
                r = portfolio.run(trades, 100_000, 0.01)
                longest, _ = underwater_stats(r["curve"])
                rows[(symbol, strat)] = (r, longest, bh)
        finally:
            backtest.FLAT_FEE_RATE = None
            sizing.FRACTIONAL = False
        print(f"  assets: {symbol} done")
    return rows


def main() -> None:
    stock_tables, headline = run_stock_portfolios()
    tf = run_timeframes()
    assets = run_assets()
    nifty_full = bh_stats(frames.daily("NIFTY 50"))
    nifty_2015 = bh_stats(frames.daily("NIFTY 50").query("ts >= '2015-02-02'")
                          .reset_index(drop=True))

    book = Workbook()
    book.remove(book.active)

    # ---------------------------------------------------------- read me ----
    sheet = book.create_sheet("Read Me")
    sheet.column_dimensions["A"].width = 110
    row = explain(sheet, 1, "WHAT THIS FILE IS", [
        "Every kitelab result in one workbook, rebuilt 2026-08-28 under the CLASS "
        "CONVENTION: the 20-EMA stack checks everything -- the stop included -- at "
        "bar CLOSES only, because the class backtests manually on end-of-bar data. "
        "The Breakout strategy keeps live intrabar stops (its buy-stop entry is "
        "inherently an intrabar order).",
        "Money conventions: stock portfolios are NET of Zerodha delivery charges and "
        "start with Rs2,50,000 unless a table says otherwise; the Timeframes sheet "
        "is GROSS (matching the class practice files); assets use flat per-side "
        "fees (0.10% crypto, 0.05% others) on 100,000 currency units.",
    ])
    row = explain(sheet, row, "HOW TO REGENERATE", [
        "This whole file: python -m scripts.master_report (signal caches in "
        "data/signal_cache feed it; delete them to force a full rebuild).",
        "Deeper detail lives one script away: scripts.tf_compare (full trade lists "
        "per timeframe), scripts.dd_proof (the formula-audited drawdown proof for "
        "the teacher), scripts.drawdown_report, scripts.full_report (the original "
        "25-minute deep-dive), scripts.showcase (practice-format trade sheets).",
        "The separate report files that used to fill this folder were removed on "
        "2026-08-28; every one of them is regenerable from those scripts. The BTC "
        "Validation sheet is the one exception -- its source upload no longer "
        "exists, so it is preserved here.",
    ])
    explain(sheet, row, "HONEST CAVEATS (apply everywhere)", [
        "Survivorship bias: the universe is TODAY'S surviving stocks -- worst for "
        "2008-2009 windows and buy-and-hold comparisons. No slippage. Kite has no "
        "intraday data before ~2015, so Breakout and hourly variants start there. "
        "MCX futures are continuous series with rollover jumps. Close-only stops "
        "mean real losses can exceed the planned 1R (nothing protects you inside "
        "the bar) -- worst singles ran 2-2.5x the planned risk.",
    ])

    # ------------------------------------------------- master comparison ----
    sheet = book.create_sheet("Master Comparison")
    row = explain(sheet, 1, "EVERYTHING, SIDE BY SIDE", [
        "Three kinds of rows, so read the group headers: portfolio accounts "
        "(compounding, net of fees), timeframe variants (fixed capital, gross), "
        "and one-asset accounts vs simply holding the asset.",
    ], span=9)
    p_headers = ["Account", "Window", "Trades", "CAGR %", "True Max DD %",
                 "Longest Underwater (yrs)", "Final Equity"]
    p_rows = []
    for name, r in headline.items():
        longest, _ = underwater_stats(r["curve"])
        first = pd.Timestamp(min(t["entry_ts"] for t in r["taken"])).year
        p_rows.append([name, f"{first}-2026", len(r["taken"]), r["cagr_pct"],
                       r["max_drawdown_pct"], longest / 365.25, r["final"]])
    p_rows.append(["NIFTY 50 buy-and-hold", "2006-2026", "-", nifty_full["cagr"],
                   nifty_full["maxdd"], nifty_full["longest_uw"] / 365.25, "-"])
    p_rows.append(["NIFTY 50 buy-and-hold", "2015-2026", "-", nifty_2015["cagr"],
                   nifty_2015["maxdd"], nifty_2015["longest_uw"] / 365.25, "-"])
    p_rows.append(["Fixed deposit (reference)", "-", "-", 7.0, 0.0, 0.0, "-"])
    row = write_table(sheet, row, "One account, 199 NSE stocks (Rs2,50,000) + benchmarks",
                      p_headers, [30, 11, 8, 8, 12, 13, 12], p_rows,
                      ["@", "@", "0", "0.0", "0.0", "0.0", "#,##0"]) + 1
    t_headers = ["Stack (5 assigned stocks)", "Trades", "Win %", "Profit Factor",
                 "Expectancy / trade", "Gross P&L"]
    t_rows = [[label.split(" -- ")[0], s["trades"], s["win_rate"],
               round(s["profit_factor"], 2), round(s["expectancy"]), round(s["gross"])]
              for label, (_, s) in tf.items()]
    row = write_table(sheet, row, "Timeframe variants (gross, common window per stock)",
                      t_headers, [26, 8, 8, 11, 13, 12], t_rows,
                      ["@", "0", "0.0%", "0.00", "#,##0", "#,##0"]) + 1
    a_headers = ["Asset", "Strategy", "CAGR %", "True Max DD %",
                 "Buy&Hold CAGR %", "Buy&Hold Max DD %"]
    a_rows = [[sym, strat, r["cagr_pct"], r["max_drawdown_pct"], bh["cagr"], bh["maxdd"]]
              for (sym, strat), (r, _, bh) in assets.items()]
    write_table(sheet, row, "Six class instruments (100,000 units, 1% risk)",
                a_headers, [13, 10, 8, 12, 12, 13], a_rows,
                ["@", "@", "0.0", "0.0", "0.0", "0.0"])

    # --------------------------------------------------------- 199 stocks ----
    sheet = book.create_sheet("199 Stocks")
    headers = ["Universe", "Risk %", "CAGR %", "True Max DD %",
               "Longest Underwater (yrs)", "Underwater Now (yrs)", "Trades", "Final Equity"]
    widths = [14, 8, 9, 12, 14, 12, 9, 13]
    fmts = ["@", "0.00", "0.0", "0.0", "0.0", "0.0", "#,##0", "#,##0"]
    row = explain(sheet, 1, "ONE ACCOUNT, MANY STOCKS", [
        f"Rs{CAPITAL:,} start, risk per trade as shown, one position per stock, "
        "signals skipped when cash runs out, NET of Zerodha delivery charges. "
        "Compare each 0.25% row with its 1% row (lower risk = shallower drawdown, "
        "and MORE profit wherever signals are plentiful), and the 150-holdout "
        "block (stocks no parameter ever saw) against in-sample.",
        f"Benchmarks: NIFTY 50 buy-and-hold 2006-2026 CAGR {nifty_full['cagr']:.1f}% "
        f"/ max DD {nifty_full['maxdd']:.1f}%; 2015-2026 {nifty_2015['cagr']:.1f}% / "
        f"{nifty_2015['maxdd']:.1f}%. A fixed deposit: ~7% with zero drawdown.",
    ], span=len(headers))
    for label, rows in stock_tables.items():
        row = write_table(sheet, row, label, headers, widths, rows, fmts) + 1

    # --------------------------------------------------------- timeframes ----
    sheet = book.create_sheet("Timeframes")
    headers = ["Stock", "Trades", "Win %", "Avg Win", "Avg Loss",
               "Expectancy / trade", "Profit Factor", "Gross P&L"]
    widths = [10, 8, 8, 11, 11, 13, 11, 12]
    fmts = ["@", "0", "0.0%", "#,##0", "#,##0", "#,##0", "0.00", "#,##0"]
    row = explain(sheet, 1, "SAME RULE, THREE TIMEFRAME STACKS", [
        "The 20-EMA stack shifted one timeframe up (Q/M/W) and one down (W/D/H) "
        "from the class strategy (M/W/D). GROSS, fixed Rs1,00,000 capital, 1% "
        "risk, common window per stock (hourly data exists only from 2015). "
        "Verdict: each trade gets better as you go UP (best PF) but total profit "
        "grows as you go DOWN (more trades) -- before charges, which multiply "
        "with trade count. Full trade lists: python -m scripts.tf_compare.",
    ], span=len(headers))
    for label, (rows, _) in tf.items():
        row = write_table(sheet, row, label, headers, widths, rows, fmts) + 1

    # ---------------------------------------------------- futures & assets ----
    sheet = book.create_sheet("Futures & Assets")
    headers = ["Asset", "Strategy", "CAGR %", "True Max DD %",
               "Longest Underwater (yrs)", "Buy&Hold CAGR %", "Buy&Hold Max DD %",
               "B&H Longest UW (yrs)"]
    widths = [13, 10, 9, 12, 14, 12, 13, 13]
    fmts = ["@", "@", "0.0", "0.0", "0.0", "0.0", "0.0", "0.0"]
    rows = [[sym, strat, r["cagr_pct"], r["max_drawdown_pct"], longest / 365.25,
             bh["cagr"], bh["maxdd"], bh["longest_uw"] / 365.25]
            for (sym, strat), (r, longest, bh) in assets.items()]
    row = explain(sheet, 1, "THE SIX CLASS INSTRUMENTS", [
        "One account per asset: 100,000 currency units, 1% risk, fractional "
        "sizing, flat per-side fees (0.10% crypto, 0.05% others). Bitcoin is in "
        "US dollars. GOLD/SILVER/CRUDEOIL are MCX continuous futures: daily data "
        "only, small rollover jumps, indices are not directly tradable. The "
        "strategies trade away return for MUCH shallower drawdowns than holding "
        "the asset raw -- on every single instrument.",
    ], span=len(headers))
    write_table(sheet, row, "Strategy accounts vs buy-and-hold",
                headers, widths, rows, fmts)

    # ----------------------------------------------------------- drawdowns ----
    sheet = book.create_sheet("Drawdowns")
    row = explain(sheet, 1, "WHAT DRAWDOWN IS AND HOW IT IS MEASURED HERE", [
        "Drawdown = how far the ACCOUNT sits below its own high-water mark. "
        "Measured properly: equity re-priced EVERY trading day (cash + shares x "
        "that day's close -- open positions at market, never at cost), each dip "
        "divided by the peak standing AT THAT MOMENT, never the final peak.",
        "Worked example: Rs100 buys 1 share at 100; closes 100, 110, 70, sell 80. "
        "Equity 100, 110, 70, 80 -> peak 110, worst day 70, max drawdown "
        "(70-110)/110 = -36%. Highest-minus-lowest is NOT a drawdown unless the "
        "peak came FIRST, and a settlement-only view would only ever see -20%.",
        "The formula-audited proof file for the teacher (every day's equity as "
        "live Excel formulas, worst-day position x-ray at checkable prices): "
        "python -m scripts.dd_proof.",
    ], span=6)
    ep_headers = ["Peak", "Trough", "Depth %", "Recovered", "Length (yrs)"]
    for label, r in headline.items():
        rows = [[str(e["peak_day"].date()), str(e["trough_day"].date()), e["depth_pct"],
                 str(e["recovered"].date()) if e["recovered"] is not None else "not yet",
                 e["days"] / 365.25] for e in episodes(r["curve"])]
        row = write_table(sheet, row, f"{label}: five deepest dips", ep_headers,
                          [13, 13, 10, 13, 11], rows, ["@", "@", "0.0", "@", "0.0"]) + 1
    col, anchor = 8, row + 1
    for label, r in list(headline.items())[:2]:
        dates, depth = underwater_series(r["curve"])
        sheet.cell(row=1, column=col, value="date")
        sheet.cell(row=1, column=col + 1, value=label)
        for i, (d, v) in enumerate(zip(dates, depth), start=2):
            sheet.cell(row=i, column=col, value=d).number_format = "yyyy-mm-dd"
            sheet.cell(row=i, column=col + 1, value=round(v, 2))
        chart = LineChart()
        chart.title = f"{label}: % below previous equity peak"
        chart.height, chart.width = 9, 22
        chart.add_data(Reference(sheet, min_col=col + 1, min_row=1, max_row=1 + len(dates)),
                       titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=col, min_row=2, max_row=1 + len(dates)))
        sheet.add_chart(chart, f"A{anchor}")
        sheet.column_dimensions[get_column_letter(col)].hidden = True
        sheet.column_dimensions[get_column_letter(col + 1)].hidden = True
        col += 2
        anchor += 19

    # ------------------------------------------------------ BTC validation ----
    if BTC_VALIDATED.exists():
        src = load_workbook(BTC_VALIDATED, data_only=False).worksheets[0]
        sheet = book.create_sheet("BTC Validation")
        for r in src.iter_rows():
            for c in r:
                if c.value is not None:
                    tgt = sheet.cell(row=c.row, column=c.column, value=c.value)
                    tgt.number_format = c.number_format
        for letter, dim in src.column_dimensions.items():
            if dim.width:
                sheet.column_dimensions[letter].width = dim.width
        sheet["A1"].alignment = src["A1"].alignment.copy()
        print("  BTC validation sheet preserved")

    target = report.save(book, "Kitelab Master Report.xlsx")
    print(f"\n  written: {target}")


if __name__ == "__main__":
    main()
